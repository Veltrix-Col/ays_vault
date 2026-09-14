"""
Genera el consolidado del ramo Movilidad a partir de cotizaciones canónicas.

Hojas: Resumen, Coberturas, Tasas, Flota_Analisis, Clausulas, Trazabilidad.
Convenciones de validación (modo CLI):
  - relleno AMARILLO  = celda que el analista debe revisar/completar (dato no encontrado, confianza < 0.7 o pendiente de LLM)
  - fuente AZUL       = puntaje/peso editable (inputs del scoring)
  - fórmulas negras   = no tocar (SUMPRODUCT, MIN, variaciones)
La hoja Trazabilidad lista cada dato con archivo/hoja/celda de origen, origen (determinista/llm) y confianza.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from gestor_cotizaciones.core.normalize import texto_para_celda
from gestor_cotizaciones.core.schema import Cotizacion, Origen

ARIAL = "Arial"
F_TIT = Font(name=ARIAL, bold=True, size=14, color="FFFFFF")
F_H = Font(name=ARIAL, bold=True, size=10, color="FFFFFF")
F_B = Font(name=ARIAL, bold=True, size=10)
F_N = Font(name=ARIAL, size=10)
F_IN = Font(name=ARIAL, size=10, color="0000FF")
FILL_T = PatternFill("solid", fgColor="1F3864")
FILL_H = PatternFill("solid", fgColor="2F5597")
FILL_G = PatternFill("solid", fgColor="D9E1F2")
FILL_REV = PatternFill("solid", fgColor="FFFF00")
FILL_ACT = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")
BORDE = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
COP = '#,##0'
PCT = '0.0%'
UMBRAL_REV = 0.7


def _hdr(ws, tomador, nit, subtitulo, actual, desde, hasta, ramo_nombre):
    ws["B2"], ws["C2"] = "TOMADOR", tomador
    ws["B3"], ws["C3"] = "NIT", nit
    ws["B4"] = ramo_nombre
    ws["B5"] = subtitulo
    ws["B6"] = f"Renovación {desde.year if desde else ''}"
    ws["B8"], ws["C8"] = "Compañía Actual", actual
    ws["B9"] = "Vigencia"
    ws["B10"], ws["C10"] = "Desde", desde
    ws["B11"], ws["C11"] = "Hasta", hasta
    for r in range(2, 12):
        ws[f"B{r}"].font = F_B
        ws[f"C{r}"].font = F_N
    for r in (4, 5, 6):
        ws[f"B{r}"].font = F_TIT; ws[f"B{r}"].fill = FILL_T
        ws.merge_cells(f"B{r}:F{r}")
    ws["C10"].number_format = ws["C11"].number_format = "yyyy-mm-dd"
    ws.column_dimensions["A"].width = 2


def _titulo(ws, row, texto, ncols=8, col=2):
    c = ws.cell(row=row, column=col, value=texto)
    c.font = F_H; c.fill = FILL_H
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + ncols - 1)


def _celda(ws, r, c, v, font=F_N, fill=None, fmt=None, align=None, border=True):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font
    if fill: cell.fill = fill
    if fmt: cell.number_format = fmt
    if align: cell.alignment = align
    if border: cell.border = BORDE
    return cell


class RenderMovilidad:
    def __init__(self, cat: dict, cots: list[Cotizacion], analisis: dict[str, Any], placas: list[str] | None = None):
        self.cat = cat
        self.cots = cots
        self.an = analisis or {}
        self.actual = next((c for c in cots if c.es_actual), None)
        # flota comparable: placas de la compañía actual (o unión) — filtrable desde CLI
        if placas:
            self.placas = [p.upper() for p in placas]
        elif self.actual:
            self.placas = [v.placa for v in self.actual.vehiculos]
        else:
            self.placas = sorted({v.placa for c in cots for v in c.vehiculos})
        self.tomador = (self.actual or cots[0]).tomador
        self.nit = (self.actual or cots[0]).nit
        self.desde = next((c.vigencia_desde for c in cots if c.vigencia_desde), None)
        self.hasta = next((c.vigencia_hasta for c in cots if c.vigencia_hasta), None)
        self.traza: list[list[Any]] = []

    # =====================================================================
    def generar(self, path: str | Path) -> Path:
        wb = Workbook()
        wb.remove(wb.active)
        self._resumen(wb.create_sheet("Resumen"))
        self._coberturas(wb.create_sheet("Coberturas"))
        self._tasas(wb.create_sheet("Tasas"))
        self._flota(wb.create_sheet("Flota_Analisis"))
        self._clausulas(wb.create_sheet("Clausulas"))
        self._trazabilidad(wb.create_sheet("Trazabilidad"))
        # el Resumen referencia totales de Flota_Analisis; se escribe de último para conocer filas
        self._resumen_ejecutivo(wb["Resumen"])
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        return Path(path)

    # ---------------------------------------------------------------- helpers
    def _hdr_std(self, ws, sub):
        _hdr(ws, self.tomador, self.nit, sub, self.actual.aseguradora_nombre if self.actual else "", self.desde, self.hasta,
             self.cat["nombre"])

    def _prima_flota(self, c: Cotizacion, anterior=False) -> float:
        d = c.por_placa()
        return float(sum(((d[p].prima_anterior_con_iva if anterior else d[p].prima_con_iva) or 0) for p in self.placas if p in d))

    # ---------------------------------------------------------------- Resumen
    def _resumen(self, ws):
        self._hdr_std(ws, "RESUMEN")
        ws.column_dimensions["B"].width = 38
        for col in "CDEFGHIJ":
            ws.column_dimensions[col].width = 18

    def _resumen_ejecutivo(self, ws):
        cat = self.cat["scoring"]
        r = 13
        _titulo(ws, r, "SCORING Y RESUMEN COMPARATIVO", 8); r += 2
        # ---- resumen ejecutivo
        _titulo(ws, r, "RESUMEN EJECUTIVO", 8); r += 1
        heads = ["Aseguradora", "Pago Anual", "Variación vs actual", "Posición Precio", "Tipo Tarifa", "Deducible PP Livianos", "Vehículo Reemplazo"]
        for j, h in enumerate(heads, start=2):
            _celda(ws, r, j, h, F_B, FILL_G, align=CENTER)
        r += 1
        prima_ant = self._prima_flota(self.actual, anterior=True) if self.actual else None
        fila_ant = r - 1  # referencia de la fila de encabezado para fórmulas
        primas = []
        for c in self.cots:
            primas.append((c, self._prima_flota(c)))
        orden = sorted(primas, key=lambda x: x[1])
        r0 = r
        for c, prima in primas:
            _celda(ws, r, 2, f"{c.aseguradora_nombre}{' (Actual)' if c.es_actual else ''}", F_B, FILL_ACT if c.es_actual else None)
            _celda(ws, r, 3, f"=Flota_Analisis!{self._col_total[c.aseguradora]}{self._fila_total}", fmt=COP)
            if prima_ant:
                _celda(ws, r, 4, f"=C{r}/Flota_Analisis!{self._col_prima_ant}{self._fila_total}-1", fmt=PCT)
            else:
                _celda(ws, r, 4, None, fill=FILL_REV)
            pos = [x[0].aseguradora for x in orden].index(c.aseguradora) + 1
            _celda(ws, r, 5, f"=RANK(C{r},$C${r0}:$C${r0 + len(primas) - 1},1)")
            _celda(ws, r, 6, c.tipo_tarifa or None, fill=None if c.tipo_tarifa else FILL_REV)
            _celda(ws, r, 7, self._ded_pp_livianos(c), align=WRAP)
            _celda(ws, r, 8, self._texto_item(c, "veh_reemplazo", "livianos"), align=WRAP)
            r += 1
        r += 2
        # ---- scoring por categoría × plan
        _titulo(ws, r, "SCORING DE COBERTURAS POR PLAN (puntajes en azul = propuesta editable; pesos editables)", 8); r += 1
        planes = [(c, p) for c in self.cots for p in c.planes if p.segmento in ("livianos", "pesados", "motos")]
        _celda(ws, r, 2, "Categoría de Cobertura", F_B, FILL_G); _celda(ws, r, 3, "Peso (%)", F_B, FILL_G)
        for j, (c, p) in enumerate(planes, start=4):
            _celda(ws, r, j, f"{c.aseguradora_nombre}\n{p.nombre}", F_B, FILL_G, align=CENTER)
        _celda(ws, r, 4 + len(planes), "Mejor Opción", F_B, FILL_G)
        ws.row_dimensions[r].height = 45
        r += 1
        r_cat0 = r
        sc = {(s["plan_id"], s["categoria_id"]): s for s in self.an.get("scoring_coberturas", [])}
        for categ in cat["categorias"]:
            _celda(ws, r, 2, categ["nombre"], F_B)
            _celda(ws, r, 3, categ["peso"], F_IN, fmt="0%")
            for j, (c, p) in enumerate(planes, start=4):
                s = sc.get((p.id, categ["id"]))
                _celda(ws, r, j, s["puntaje"] if s else None, F_IN, None if s else FILL_REV)
                if s:
                    ws.cell(row=r, column=j).comment = None
            c1, c2 = L(4), L(3 + len(planes))
            _celda(ws, r, 4 + len(planes), f'=IFERROR(INDEX(${c1}${r_cat0 - 1}:${c2}${r_cat0 - 1},MATCH(MAX({c1}{r}:{c2}{r}),{c1}{r}:{c2}{r},0)),"")', align=WRAP)
            r += 1
        _celda(ws, r, 2, "SCORE TOTAL PONDERADO", F_B, FILL_G)
        _celda(ws, r, 3, f"=SUM(C{r_cat0}:C{r - 1})", F_B, FILL_G, fmt="0%")
        for j in range(4, 4 + len(planes)):
            cl = L(j)
            _celda(ws, r, j, f"=SUMPRODUCT({cl}{r_cat0}:{cl}{r - 1},$C${r_cat0}:$C${r - 1})", F_B, FILL_G, fmt="0.00")
        c1, c2 = L(4), L(3 + len(planes))
        _celda(ws, r, 4 + len(planes), f'=IFERROR(INDEX(${c1}${r_cat0 - 1}:${c2}${r_cat0 - 1},MATCH(MAX({c1}{r}:{c2}{r}),{c1}{r}:{c2}{r},0)),"")', F_B, FILL_G, align=WRAP)
        r += 3
        # ---- scoring por tipo de vehículo
        _titulo(ws, r, "SCORING POR TIPO DE VEHÍCULO", 8); r += 1
        _celda(ws, r, 2, "Tipo de Vehículo", F_B, FILL_G)
        for j, c in enumerate(self.cots, start=3):
            _celda(ws, r, j, c.aseguradora_nombre, F_B, FILL_G, align=CENTER)
        _celda(ws, r, 3 + len(self.cots), "Mejor Opción", F_B, FILL_G); _celda(ws, r, 4 + len(self.cots), "Observaciones", F_B, FILL_G)
        r += 1
        tipos = ["Automóviles/Camperos/Pickups", "Vehículos Blindados", "Vehículos Pesados", "Motocicletas > 250cc", "Motocicletas < 250cc"]
        st = {(s["tipo"].lower()[:12], s["aseguradora"]): s for s in self.an.get("scoring_tipo_vehiculo", [])}
        r_t0 = r
        for t in tipos:
            _celda(ws, r, 2, t, F_B)
            obs = []
            for j, c in enumerate(self.cots, start=3):
                s = st.get((t.lower()[:12], c.aseguradora))
                _celda(ws, r, j, s["puntaje"] if s else None, F_IN, None if s else FILL_REV)
                if s and s.get("observacion"): obs.append(f"{c.aseguradora_nombre}: {s['observacion']}")
            c1, c2 = L(3), L(2 + len(self.cots))
            _celda(ws, r, 3 + len(self.cots), f'=IFERROR(INDEX(${c1}${r_t0 - 1}:${c2}${r_t0 - 1},MATCH(MAX({c1}{r}:{c2}{r}),{c1}{r}:{c2}{r},0)),"")')
            _celda(ws, r, 4 + len(self.cots), " | ".join(obs) or None, align=WRAP, fill=None if obs else FILL_REV)
            r += 1
        _celda(ws, r, 2, "PROMEDIO FLOTA", F_B, FILL_G)
        for j in range(3, 3 + len(self.cots)):
            _celda(ws, r, j, f"=IFERROR(AVERAGE({L(j)}{r_t0}:{L(j)}{r - 1}),\"\")", F_B, FILL_G, fmt="0.00")
        c1, c2 = L(3), L(2 + len(self.cots))
        _celda(ws, r, 3 + len(self.cots), f'=IFERROR(INDEX(${c1}${r_t0 - 1}:${c2}${r_t0 - 1},MATCH(MAX({c1}{r}:{c2}{r}),{c1}{r}:{c2}{r},0)),"")', F_B, FILL_G)
        r += 3
        # ---- ventajas / desventajas
        _titulo(ws, r, "VENTAJAS Y DESVENTAJAS POR ASEGURADORA", 2 * len(self.cots)); r += 1
        vd = {v["aseguradora"]: v for v in self.an.get("ventajas_desventajas", [])}
        for j, c in enumerate(self.cots):
            col = 2 + 2 * j
            _celda(ws, r, col, f"{c.aseguradora_nombre}{' (Compañía Actual)' if c.es_actual else ''}", F_B, FILL_G, align=CENTER)
            ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + 1)
            _celda(ws, r + 1, col, "VENTAJAS", F_B, FILL_ACT); _celda(ws, r + 1, col + 1, "DESVENTAJAS", F_B, PatternFill("solid", fgColor="FCE4D6"))
        r += 2
        n = max([max(len(v["ventajas"]), len(v["desventajas"])) for v in vd.values()] + [5])
        for k in range(n):
            for j, c in enumerate(self.cots):
                v = vd.get(c.aseguradora)
                col = 2 + 2 * j
                _celda(ws, r + k, col, v["ventajas"][k] if v and k < len(v["ventajas"]) else None, align=WRAP, fill=None if v else FILL_REV)
                _celda(ws, r + k, col + 1, v["desventajas"][k] if v and k < len(v["desventajas"]) else None, align=WRAP, fill=None if v else FILL_REV)
        r += n + 2
        # ---- matriz de decisión
        _titulo(ws, r, "MATRIZ DE DECISIÓN FINAL", 8); r += 1
        _celda(ws, r, 2, "Criterio", F_B, FILL_G); _celda(ws, r, 3, "Peso", F_B, FILL_G)
        for j, c in enumerate(self.cots, start=4):
            _celda(ws, r, j, c.aseguradora_nombre, F_B, FILL_G, align=CENTER)
        r += 1
        md = {(m["aseguradora"], m["criterio_id"]): m["puntaje"] for m in self.an.get("matriz_decision", [])}
        r_m0 = r
        for crit in cat["matriz_decision"]:
            _celda(ws, r, 2, crit["nombre"], F_B); _celda(ws, r, 3, crit["peso"], F_IN, fmt="0%")
            for j, c in enumerate(self.cots, start=4):
                p = md.get((c.aseguradora, crit["id"]))
                _celda(ws, r, j, p, F_IN, None if p is not None else FILL_REV)
            r += 1
        _celda(ws, r, 2, "SCORE FINAL", F_B, FILL_G); _celda(ws, r, 3, f"=SUM(C{r_m0}:C{r - 1})", F_B, FILL_G, fmt="0%")
        for j in range(4, 4 + len(self.cots)):
            _celda(ws, r, j, f"=SUMPRODUCT({L(j)}{r_m0}:{L(j)}{r - 1},$C${r_m0}:$C${r - 1})", F_B, FILL_G, fmt="0.00")
        c1, c2 = L(4), L(3 + len(self.cots))
        _celda(ws, r, 4 + len(self.cots), f'=IFERROR("GANADOR: "&INDEX(${c1}${r_m0 - 1}:${c2}${r_m0 - 1},MATCH(MAX({c1}{r}:{c2}{r}),{c1}{r}:{c2}{r},0)),"")', F_B, FILL_G)
        r += 3
        # ---- conclusiones
        _titulo(ws, r, "CONCLUSIONES Y RECOMENDACIÓN (propuesta del modelo, validar)", 8); r += 1
        for j, h in enumerate(["Perfil", "Recomendación", "Justificación"], start=2):
            _celda(ws, r, j, h, F_B, FILL_G)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=9)
        r += 1
        recs = self.an.get("recomendaciones", [])
        for k in range(max(len(recs), 5)):
            rec = recs[k] if k < len(recs) else None
            _celda(ws, r, 2, rec["perfil"] if rec else None, F_B, None if rec else FILL_REV)
            _celda(ws, r, 3, rec["recomendacion"] if rec else None, align=WRAP, fill=None if rec else FILL_REV)
            _celda(ws, r, 4, rec["justificacion"] if rec else None, align=WRAP, fill=None if rec else FILL_REV)
            ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=9)
            r += 1
        r += 1
        c = _celda(ws, r, 2, self.an.get("conclusion") or None, align=WRAP, fill=None if self.an.get("conclusion") else FILL_REV)
        ws.merge_cells(start_row=r, start_column=2, end_row=r + 3, end_column=9)
        r += 5
        _celda(ws, r, 2, "Leyenda: amarillo = pendiente de validar/completar por el analista · azul = puntaje/peso editable · negro = fórmula", Font(name=ARIAL, italic=True, size=9), border=False)

    def _ded_pp_livianos(self, c: Cotizacion):
        for p in c.planes:
            if p.segmento == "livianos":
                it = c.item("ppd", p.id, "deducible")
                if it:
                    return texto_para_celda(it.valor)
        return None

    def _texto_item(self, c: Cotizacion, cid: str, segmento: str):
        for p in c.planes:
            if p.segmento == segmento:
                it = c.item(cid, p.id)
                if it:
                    return texto_para_celda(it.valor)
        return None

    # ---------------------------------------------------------------- Coberturas
    def _coberturas(self, ws):
        self._hdr_std(ws, "INFORMACIÓN BÁSICA")
        ws.column_dimensions["B"].width = 48
        orden_seg = [s["id"] for s in self.cat["segmentos"]]
        planes = [(c, p) for c in self.cots for p in sorted(c.planes, key=lambda p: orden_seg.index(p.segmento) if p.segmento in orden_seg else 99)]
        col0 = 3
        for j, (c, p) in enumerate(planes, start=col0):
            ws.column_dimensions[L(j)].width = 22

        def encabezado(r, titulo):
            # fila aseguradora (merge por aseguradora), fila segmento, fila plan
            j = col0
            for c in self.cots:
                n = sum(1 for cc, _ in planes if cc is c)
                _celda(ws, r, j, c.aseguradora_nombre, F_H, FILL_H, align=CENTER)
                ws.merge_cells(start_row=r, start_column=j, end_row=r, end_column=j + n - 1)
                j += n
            for j, (c, p) in enumerate(planes, start=col0):
                seg = next(s["nombre"] for s in self.cat["segmentos"] if s["id"] == p.segmento)
                _celda(ws, r + 1, j, seg, F_B, FILL_G, align=CENTER)
                _celda(ws, r + 2, j, p.nombre, F_B, FILL_G, align=CENTER)
            _celda(ws, r + 2, 2, titulo, F_H, FILL_H)
            ws.row_dimensions[r + 1].height = 30; ws.row_dimensions[r + 2].height = 30
            return r + 3

        r = encabezado(13, "COBERTURAS")
        for g in self.cat["coberturas"]:
            _celda(ws, r, 2, g["grupo"], F_B, FILL_G)
            for j in range(col0, col0 + len(planes)):
                _celda(ws, r, j, None, fill=FILL_G)
            r += 1
            for item in g["items"]:
                _celda(ws, r, 2, item["nombre"], F_N, align=WRAP)
                for j, (c, p) in enumerate(planes, start=col0):
                    self._celda_item(ws, r, j, c, p.id, item["id"], "cobertura")
                r += 1
        r += 1
        r = encabezado(r, "DEDUCIBLES")
        for d in self.cat["deducibles"]:
            _celda(ws, r, 2, d["nombre"], F_N, align=WRAP)
            for j, (c, p) in enumerate(planes, start=col0):
                self._celda_item(ws, r, j, c, p.id, d["id"], "deducible")
            r += 1
        r += 1
        # generales (una columna por aseguradora: se escribe en la primera columna de cada una)
        gen = {
            "modalidad_pago": lambda c: c.modalidad_pago, "descuento_flota": lambda c: c.descuento_flota,
            "guia_fasecolda": lambda c: c.guia_fasecolda, "condicionados": lambda c: "\n".join(c.condicionados) or None,
            "notas": lambda c: "\n".join(c.notas) or None,
        }
        for g in self.cat["generales"]:
            _celda(ws, r, 2, g["nombre"], F_B)
            j = col0
            for c in self.cots:
                n = sum(1 for cc, _ in planes if cc is c)
                v = gen[g["id"]](c)
                _celda(ws, r, j, v, align=WRAP, fill=None if v is not None else FILL_REV,
                       fmt="0%" if g["id"] == "descuento_flota" else None)
                ws.merge_cells(start_row=r, start_column=j, end_row=r, end_column=j + n - 1)
                j += n
            r += 1
        ws.freeze_panes = "C16"

    def _celda_item(self, ws, r, j, c: Cotizacion, plan_id, cid, tipo):
        it = c.item(cid, plan_id, tipo)
        if it is None:
            _celda(ws, r, j, None, fill=FILL_REV, align=WRAP)
            return
        v = it.valor
        val = texto_para_celda(v)
        fill = FILL_REV if (v.origen == Origen.LLM and v.confianza < UMBRAL_REV) else None
        cell = _celda(ws, r, j, val, fill=fill, align=WRAP, fmt=COP if v.unidad == "COP" else None)
        f = v.fuente
        self.traza.append([c.aseguradora_nombre, "Coberturas", f"{L(j)}{r}", plan_id, cid, tipo, val,
                           f.archivo if f else "", f.hoja if f else "", f.celda if f else "", v.origen.value, v.confianza, v.nota or ""])

    # ---------------------------------------------------------------- Tasas
    def _tasas(self, ws):
        self._hdr_std(ws, "TASAS")
        ws.column_dimensions["B"].width = 34
        r = 13
        for c in self.cots:
            _titulo(ws, r, f"{c.aseguradora_nombre}  ({c.tipo_tarifa or ''})", 8); r += 1
            if not c.tasas:
                _celda(ws, r, 2, "Sin tabla de tasas en la cotización (tarifa individual por vehículo: ver Flota_Analisis)", fill=FILL_REV)
                r += 3
                continue
            segs = list(dict.fromkeys(t.segmento for t in c.tasas))
            rangos = list(dict.fromkeys(t.rango for t in c.tasas))
            _celda(ws, r, 2, "Rango", F_B, FILL_G)
            for j, s in enumerate(segs, start=3):
                _celda(ws, r, j, s, F_B, FILL_G, align=CENTER); ws.column_dimensions[L(j)].width = max(ws.column_dimensions[L(j)].width or 0, 16)
            ws.row_dimensions[r].height = 42
            r += 1
            idx = {(t.segmento, t.rango): t.tasa for t in c.tasas}
            for rg in rangos:
                _celda(ws, r, 2, rg, F_B)
                for j, s in enumerate(segs, start=3):
                    _celda(ws, r, j, idx.get((s, rg)), fmt="0.000%")
                r += 1
            # adicionales por vehículo (asistencia, reemplazo) si existen en extra
            extras = {}
            for v in c.vehiculos:
                for k in ("asistencia", "veh_reemplazo", "pac"):
                    if v.extra.get(k):
                        extras.setdefault(f"{k} ({v.clase})", v.extra[k])
            if extras:
                r += 1
                _celda(ws, r, 2, "Cobros adicionales por vehículo (COP, sin IVA)", F_B, FILL_G); r += 1
                for k, val in extras.items():
                    _celda(ws, r, 2, k); _celda(ws, r, 3, val, fmt=COP); r += 1
            r += 2

    # ---------------------------------------------------------------- Flota
    def _flota(self, ws):
        self._hdr_std(ws, "LIQUIDACIÓN RIESGOS")
        base = self.actual or self.cots[0]
        d_base = base.por_placa()
        r_h = 13
        fijas = ["Placa", "Marca - Referencia", "Chasis", "Motor", "Modelo", "Ciudad", "Clase", "Código Fasecolda", "Valor asegurado actual"]
        for j, h in enumerate(fijas, start=2):
            _celda(ws, r_h, j, h, F_B, FILL_G, align=CENTER)
        j = 2 + len(fijas)
        self._col_total = {}
        cols_var_prima = []
        for c in self.cots:
            _celda(ws, r_h - 1, j, c.aseguradora_nombre, F_H, FILL_H, align=CENTER)
            if c.es_actual:
                heads = ["Valor asegurado 2026", "Valor accesorios", "Total valor asegurado", "Var % VA", "N° siniestros", "Pago anual anterior (con IVA)", "Pago anual 2026 (con IVA)", "Var % prima", "Pago mensual 2026"]
            else:
                heads = ["Valor asegurado 2026", "Valor accesorios", "Total valor asegurado", "Var % VA", "Pago anual 2026 (con IVA)", "Var % vs anterior", "Pago mensual 2026"]
            ws.merge_cells(start_row=r_h - 1, start_column=j, end_row=r_h - 1, end_column=j + len(heads) - 1)
            for k, h in enumerate(heads):
                _celda(ws, r_h, j + k, h, F_B, FILL_G, align=CENTER)
            c.__dict__["_col0"] = j
            j += len(heads)
        _celda(ws, r_h, j, "Pago anual más bajo", F_B, FILL_G, align=CENTER); _celda(ws, r_h, j + 1, "Aseguradora pago más bajo", F_B, FILL_G, align=CENTER)
        col_min, col_min_aseg = j, j + 1
        ws.row_dimensions[r_h].height = 45
        r = r_h + 1
        r0 = r
        col_va_actual = 2 + len(fijas) - 1
        for placa in self.placas:
            vb = d_base.get(placa)
            vals = [placa, vb.marca_referencia if vb else None, vb.chasis if vb else None, vb.motor if vb else None,
                    vb.modelo if vb else None, vb.ciudad if vb else None, vb.clase if vb else None,
                    vb.codigo_fasecolda if vb else None, vb.valor_asegurado if vb else None]
            for k, v in enumerate(vals, start=2):
                _celda(ws, r, k, v, fmt=COP if k == col_va_actual else None, fill=None if v is not None else FILL_REV)
            cols_prima = []
            for c in self.cots:
                v = c.por_placa().get(placa)
                j = c.__dict__["_col0"]
                A, B, C = L(j), L(j + 1), L(j + 2)
                _celda(ws, r, j, v.valor_asegurado if v else None, fmt=COP, fill=None if v else FILL_REV)
                _celda(ws, r, j + 1, v.valor_accesorios if v else None, fmt=COP, fill=None if v else FILL_REV)
                _celda(ws, r, j + 2, f"={A}{r}+{B}{r}", fmt=COP)
                _celda(ws, r, j + 3, f'=IF({L(col_va_actual)}{r}>0,{C}{r}/{L(col_va_actual)}{r}-1,"")', fmt=PCT)
                if c.es_actual:
                    _celda(ws, r, j + 4, v.numero_siniestros if v else None)
                    _celda(ws, r, j + 5, v.prima_anterior_con_iva if v else None, fmt=COP, fill=None if v else FILL_REV)
                    _celda(ws, r, j + 6, v.prima_con_iva if v else None, fmt=COP, fill=None if v else FILL_REV)
                    _celda(ws, r, j + 7, f'=IF({L(j + 5)}{r}>0,{L(j + 6)}{r}/{L(j + 5)}{r}-1,"")', fmt=PCT)
                    _celda(ws, r, j + 8, f"={L(j + 6)}{r}/12", fmt=COP)
                    self._col_prima_ant = L(j + 5)
                    cols_prima.append(L(j + 6))
                else:
                    _celda(ws, r, j + 4, v.prima_con_iva if v else None, fmt=COP, fill=None if v else FILL_REV)
                    ref_ant = f"{self._col_prima_ant}{r}" if self.actual else "0"
                    _celda(ws, r, j + 5, f'=IF({ref_ant}>0,{L(j + 4)}{r}/{ref_ant}-1,"")', fmt=PCT)
                    _celda(ws, r, j + 6, f"={L(j + 4)}{r}/12", fmt=COP)
                    cols_prima.append(L(j + 4))
            refs = ",".join(f"{cl}{r}" for cl in cols_prima)
            _celda(ws, r, col_min, f"=MIN({refs})", fmt=COP)
            nombres = [c.aseguradora_nombre for c in self.cots]
            f = "".join(f'IF({L(col_min)}{r}={cl}{r},"{n}",' for cl, n in zip(cols_prima[:-1], nombres[:-1])) + f'"{nombres[-1]}"' + ")" * (len(cols_prima) - 1)
            _celda(ws, r, col_min_aseg, f"={f}")
            r += 1
        # totales
        _celda(ws, r, 2, "TOTAL", F_B, FILL_G)
        _celda(ws, r, col_va_actual, f"=SUM({L(col_va_actual)}{r0}:{L(col_va_actual)}{r - 1})", F_B, FILL_G, fmt=COP)
        for c in self.cots:
            j = c.__dict__["_col0"]
            n = 9 if c.es_actual else 7
            for k in range(n):
                cl = L(j + k)
                hdr = ws.cell(row=r_h, column=j + k).value
                if hdr.startswith("Var %"):
                    prev = L(j + k - 1); base_col = L(col_va_actual) if "VA" in hdr else (self._col_prima_ant if self.actual else None)
                    _celda(ws, r, j + k, f'=IF({base_col}{r}>0,{prev}{r}/{base_col}{r}-1,"")' if base_col else None, F_B, FILL_G, fmt=PCT)
                elif hdr.startswith("N°"):
                    _celda(ws, r, j + k, f"=SUM({cl}{r0}:{cl}{r - 1})", F_B, FILL_G)
                else:
                    _celda(ws, r, j + k, f"=SUM({cl}{r0}:{cl}{r - 1})", F_B, FILL_G, fmt=COP)
                if hdr.startswith("Pago anual 2026"):
                    self._col_total[c.aseguradora] = cl
        _celda(ws, r, col_min, f"=SUM({L(col_min)}{r0}:{L(col_min)}{r - 1})", F_B, FILL_G, fmt=COP)
        self._fila_total = r
        # vehículos cotizados por otras aseguradoras pero fuera de la flota comparable
        extra = sorted({v.placa for c in self.cots for v in c.vehiculos} - set(self.placas))
        if extra:
            r += 2
            _celda(ws, r, 2, f"Placas cotizadas pero excluidas de la flota comparable ({len(extra)}): {', '.join(extra)}",
                   Font(name=ARIAL, italic=True, size=9), border=False)
        ws.freeze_panes = f"C{r0}"
        for j in range(2, col_min_aseg + 1):
            ws.column_dimensions[L(j)].width = 16
        ws.column_dimensions["C"].width = 34

    # ---------------------------------------------------------------- Cláusulas
    def _clausulas(self, ws):
        self._hdr_std(ws, "CLÁUSULAS, CONDICIONES Y EXCLUSIONES")
        r = 13
        for j, h in enumerate(["Aseguradora", "Categoría", "Título", "Texto", "Fuente", "Confianza"], start=2):
            _celda(ws, r, j, h, F_B, FILL_G)
        r += 1
        for c in self.cots:
            if not c.clausulas:
                _celda(ws, r, 2, c.aseguradora_nombre); _celda(ws, r, 3, "Pendiente de extracción (LLM)", fill=FILL_REV); r += 1
            for cl in sorted(c.clausulas, key=lambda x: x.categoria):
                _celda(ws, r, 2, c.aseguradora_nombre); _celda(ws, r, 3, cl.categoria); _celda(ws, r, 4, cl.titulo, align=WRAP)
                _celda(ws, r, 5, cl.texto, align=WRAP); _celda(ws, r, 6, cl.fuente.archivo if cl.fuente else None, align=WRAP)
                _celda(ws, r, 7, cl.confianza, fmt="0.00", fill=FILL_REV if cl.confianza < UMBRAL_REV else None)
                r += 1
        for col, w in zip("BCDEFG", (16, 20, 40, 90, 40, 10)):
            ws.column_dimensions[col].width = w

    # ---------------------------------------------------------------- Trazabilidad
    def _trazabilidad(self, ws):
        heads = ["Aseguradora", "Hoja", "Celda consolidado", "Plan", "Cobertura", "Tipo", "Valor", "Archivo origen", "Hoja origen", "Celda/ref origen", "Origen", "Confianza", "Nota"]
        for j, h in enumerate(heads, start=1):
            _celda(ws, 1, j, h, F_B, FILL_G)
        for i, fila in enumerate(self.traza, start=2):
            for j, v in enumerate(fila, start=1):
                _celda(ws, i, j, v, fill=FILL_REV if (fila[10] == "llm" and fila[11] < UMBRAL_REV) else None)
        for j, w in enumerate((14, 12, 10, 24, 22, 10, 40, 30, 22, 14, 12, 10, 30), start=1):
            ws.column_dimensions[L(j)].width = w
        ws.auto_filter.ref = f"A1:{L(len(heads))}{max(2, len(self.traza) + 1)}"
        ws.freeze_panes = "A2"
