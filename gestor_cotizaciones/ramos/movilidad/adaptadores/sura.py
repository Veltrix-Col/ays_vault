"""
SURA — cotización de renovación (xlsx con una hoja 'Flota Actual <póliza>' por póliza) + slip (docx).

Determinista:
  - 'Resumen Polizas' : vigencia, pólizas ↔ tomador, descuento, tabla de deducibles por plan.
  - 'Flota Actual *'  : vehículo a vehículo (incluye prima anterior → sirve como compañía actual).
LLM:
  - Slip .docx : coberturas por plan (Básico / Global / Clásico / Utilitarios / Pesados / Motos),
                 asistencias, beneficios, cláusulas y exclusiones.

Si se pasa opciones={"nit": ...} solo se toman las pólizas de ese tomador.
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from gestor_cotizaciones.core.adaptador import Adaptador, extraer_con_llm
from gestor_cotizaciones.core.ingest import bloques_docx
from gestor_cotizaciones.core.normalize import a_fecha, a_numero, compacta, deducible, limpiar
from gestor_cotizaciones.core.schema import Fuente, ItemCobertura, Plan, Valor, Vehiculo

# Columnas de la tabla de deducibles de 'Resumen Polizas' -> plan canónico
_PLANES_DEDUC = {
    "Cobertura 100%": ("sura.global", "Global con Franquicia", "livianos"),
    "Plan Básico": ("sura.basico", "Básico", "livianos"),
    "Motos Estandar (CC <= 250)": ("sura.motos_250", "Menos de 250 cc", "motos"),
    "Motos Alto Cilindraje (CC > 250)": ("sura.motos_alto", "Más de 250 cc", "motos"),
    "Motocarros": ("sura.motocarros", "Motocarros", "motos"),
    "Utilitarios y Camión Liviano": ("sura.utilitarios", "Plan Utilitarios y Camión Liviano", "pesados"),
    "Camiones Medianos; Pesados y Otros Pesados": ("sura.pesados", "Camión Mediano, Pesado y Otros Pesados", "pesados"),
    "Vehículos Blindados": ("sura.blindados", "Blindados", "blindados"),
}
# Planes adicionales que solo aparecen en el slip
_PLANES_SLIP = [
    ("sura.global_deducible", "Global con Deducible", "livianos"),
    ("sura.clasico", "Clásico Suma Fija Plus", "livianos"),
]
_DEDUC_FILAS = {"R.C.": "rc_bienes", "P.T.D.": "ptd", "P.P.D.": "ppd", "P.T.H.": "pth", "P.P.H": "pph", "P.P.H.": "pph"}


class AdaptadorSura(Adaptador):
    ramo = "movilidad"
    id = "sura"
    nombre = "SURA"
    patrones = [r"sura"]

    def extraer(self, archivos):
        cot = self.nueva(archivos)
        xlsx = next(a for a in archivos if str(a).lower().endswith(".xlsx"))
        docx = next((a for a in archivos if str(a).lower().endswith(".docx")), None)
        wb = load_workbook(xlsx, read_only=True, data_only=True)
        polizas = self._resumen(cot, wb["Resumen Polizas"], xlsx)
        nit = re.sub(r"\D", "", str(self.opciones.get("nit", "")))
        for pol, info in polizas.items():
            if nit and nit not in info["nit"]:
                continue
            hoja = f"Flota Actual {pol}"
            if hoja in wb.sheetnames:
                self._vehiculos(cot, wb[hoja], xlsx, pol)
                cot.tomador = cot.tomador or info["tomador"]
                cot.nit = cot.nit or info["nit"]
        cot.tipo_tarifa = "Tasa Individual"
        for pid, nombre, seg in _PLANES_SLIP:
            cot.planes.append(Plan(id=pid, nombre=nombre, segmento=seg, descripcion="Descrito en el slip"))
        if self.llm and docx:
            bl = bloques_docx(docx)
            extraer_con_llm(cot, bl, self.cat, self.llm, instrucciones_extra=(
                "Los deducibles de RC/PTD/PPD/PTH/PPH por plan YA fueron extraídos del Excel. Del slip extrae, POR PLAN "
                "(Básico, Global, Clásico, Utilitarios y Camión Liviano, Pesados, Motos), el límite de RC, si se otorgan "
                "PTD/PPD/PTH/PPH, gastos de transporte, vehículo de reemplazo (días), accidentes al conductor y ocupantes, "
                "cada asistencia y beneficio con su límite, cláusulas, exclusiones y condicionados vigentes. "
                "'Plan Autos Global' corresponde a sura.global y sura.global_deducible (mismas coberturas salvo el deducible)."))
        return cot

    # ------------------------------------------------------------------
    def _resumen(self, cot, ws, xlsx) -> dict[str, dict]:
        rows = list(ws.iter_rows(min_row=1, max_row=40, values_only=True))
        polizas: dict[str, dict] = {}
        for i, r in enumerate(rows, 1):
            c = compacta(r)
            t = [limpiar(x) for x in c]
            if t and t[0] == "VIGENCIA:":
                fechas = [a_fecha(x) for x in c if a_fecha(x)]
                if len(fechas) >= 2:
                    cot.vigencia_desde, cot.vigencia_hasta = fechas[0], fechas[1]
            if t and t[0] == "Póliza":
                for rr in rows[i:]:
                    cc = compacta(rr)
                    if len(cc) < 4 or not re.fullmatch(r"\d{6,}", limpiar(cc[0])):
                        break
                    tom = limpiar(cc[1])
                    m = re.search(r"NIT\s*([\d.\- ]+)", tom, re.I)
                    polizas[limpiar(cc[0])] = {"tomador": re.split(r"\s+NIT", tom, flags=re.I)[0].strip(),
                                               "nit": re.sub(r"\D", "", m.group(1)) if m else "", "modalidad": limpiar(cc[3])}
                    cot.modalidad_pago = cot.modalidad_pago or ("Mensual" if "MENSUAL" in limpiar(cc[3]).upper() else limpiar(cc[3]))
            if t and t[0] == "Descuento póliza" and i < len(rows):
                d = compacta(rows[i])
                if d:
                    cot.descuento_flota = a_numero(d[0])
            t = [limpiar(x) for x in r]   # para la tabla de deducibles se necesitan índices reales de columna
            if "Amparos" in t:
                cols = {j: _PLANES_DEDUC[limpiar(c)] for j, c in enumerate(r) if limpiar(c) in _PLANES_DEDUC}
                for pid, nombre, seg in cols.values():
                    cot.planes.append(Plan(id=pid, nombre=nombre, segmento=seg, descripcion="Columna de deducibles en 'Resumen Polizas'"))
                for k, rr in enumerate(rows[i:i + 8], start=i + 1):
                    cid = _DEDUC_FILAS.get(limpiar(rr[2]))
                    if not cid:
                        continue
                    for j, (pid, _, _) in cols.items():
                        d = deducible(rr[j]); d.fuente = Fuente(archivo=Path(xlsx).name, hoja=ws.title, celda=f"{chr(65 + j)}{k}")
                        cot.items.append(ItemCobertura(cobertura_id=cid, plan_id=pid, tipo="deducible", valor=d))
                        if cid != "rc_bienes":
                            ap = d.aplica is not False
                            cot.items.append(ItemCobertura(cobertura_id=cid, plan_id=pid, valor=Valor(
                                texto="Si" if ap else "No", aplica=ap, valor=ap, unidad="bool", fuente=d.fuente)))
                # 'Global con Deducible' comparte deducibles con Global salvo PPD/PPH = 1 SMMLV (ver slip); se deja al LLM
        return polizas

    def _vehiculos(self, cot, ws, xlsx, poliza):
        rows = list(ws.iter_rows(values_only=True))
        hdr = [limpiar(c) for c in rows[0]]
        col = {n: hdr.index(n) for n in hdr if n}
        for r in rows[1:]:
            if not r or not r[0] or limpiar(r[0]).startswith("="):
                continue
            placa = limpiar(r[col["Placa"]]).upper()
            if not re.fullmatch(r"[A-Z0-9]{5,7}", placa):
                continue
            clase, _ = self.clase_canonica(r[col["Clase"]])
            prima_ant_sin_iva = a_numero(r[col["Prima Anterior"]])
            cot.vehiculos.append(Vehiculo(
                placa=placa, codigo_fasecolda=str(r[col["Fasecolda"]]).zfill(8) if r[col["Fasecolda"]] else None,
                marca_referencia=f"{limpiar(r[col['Marca']])} {limpiar(r[col['Linea']])}".strip(),
                modelo=int(r[col["Modelo"]]) if r[col["Modelo"]] else None, clase=clase,
                ciudad=limpiar(r[col["Ciudad de Circulación"]]).split("-")[0], chasis=limpiar(r[col["Chasis"]]),
                motor=limpiar(r[col["Motor"]]), plan=limpiar(r[col["Plan"]]),
                valor_asegurado=a_numero(r[col["Valor Asegurado"]]), valor_accesorios=a_numero(r[col["Valor Accesorios"]]) or 0.0,
                valor_total=(a_numero(r[col["Valor Asegurado"]]) or 0) + (a_numero(r[col["Valor Accesorios"]]) or 0),
                prima_neta=a_numero(r[col["Prima sin IVA"]]), prima_con_iva=a_numero(r[col["Prima con IVA"]]),
                prima_anterior_con_iva=round(prima_ant_sin_iva * 1.19) if prima_ant_sin_iva else None,
                numero_siniestros=int(r[col["Numero de Siniestros"]] or 0),
                extra={"poliza": poliza, "rc_limite": limpiar(r[col["Responsabilidad Civil-Limite"]]),
                       "descuento": a_numero(r[col["Descuento"]]), "recargo": a_numero(r[col["Recargo"]])},
            ))
