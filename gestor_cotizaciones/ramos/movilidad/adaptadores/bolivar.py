"""
SEGUROS BOLÍVAR — cotización (xlsx con TablaSalida / TasasSalida / ResumenSalida) + slip (xlsx clave-valor).

Determinista:
  - ResumenSalida : fecha, tomador, tipo de tasa, guía Fasecolda, forma de pago, tabla RC/deducibles por condición.
  - TablaSalida   : vehículo a vehículo.
  - TasasSalida   : tasas por clase y rango de valor asegurado.
LLM:
  - Slip 'CONDICIONES PARTICULARES Y BENE...' + 'TENER EN CUENTA' de ResumenSalida.
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from gestor_cotizaciones.core.adaptador import Adaptador, extraer_con_llm
from gestor_cotizaciones.core.ingest import bloques_xlsx
from gestor_cotizaciones.core.normalize import a_fecha, a_numero, a_tasa, deducible, limpiar
from gestor_cotizaciones.core.schema import Fuente, ItemCobertura, Plan, Tasa, Valor, Vehiculo

# Condición (fila de ResumenSalida) -> plan canónico
_PLANES = {
    "livianos familiares hasta 20": ("bolivar.livianos", "Automóviles, Camioneta, Camperos y Pickup", "livianos"),
    "livanos familiares mayor a 20": ("bolivar.livianos_20", "Automóviles, Camioneta, Camperos y Pickup mas de 20 años", "livianos"),
    "livianos familiares mayor a 20": ("bolivar.livianos_20", "Automóviles, Camioneta, Camperos y Pickup mas de 20 años", "livianos"),
    "pesados": ("bolivar.pesados", "Pesados", "pesados"),
    "motos": ("bolivar.motos", "Motos", "motos"),
}


class AdaptadorBolivar(Adaptador):
    ramo = "movilidad"
    id = "bolivar"
    nombre = "BOLIVAR"
    patrones = [r"bolivar", r"bol[ií]var"]

    def extraer(self, archivos):
        cot = self.nueva(archivos)
        cotiz = next(a for a in archivos if "slip" not in str(a).lower())
        slip = next((a for a in archivos if "slip" in str(a).lower()), None)
        wb = load_workbook(cotiz, read_only=True, data_only=True)
        self._resumen(cot, wb["ResumenSalida"], cotiz)
        self._vehiculos(cot, wb["TablaSalida"], cotiz)
        self._tasas(cot, wb["TasasSalida"], cotiz)
        if self.llm:
            bl = bloques_xlsx(cotiz, hojas=["ResumenSalida"])
            bl = [b for b in bl if int(b.ref.split()[1]) >= 34]
            if slip:
                bl += bloques_xlsx(slip)
            extraer_con_llm(cot, bl, self.cat, self.llm, instrucciones_extra=(
                "Los límites de RC y los deducibles YA fueron extraídos. Concéntrate en PTD/PPD/PTH/PPH (si se otorgan), "
                "gastos de transporte, vehículo de reemplazo (opción premium / días), accidentes personales, asistencias, "
                "beneficios, cláusulas, condicionados vigentes y notas (abonos parciales, antigüedad máxima, etc.). "
                "Cuando el slip no distingue plan, aplica el item a 'bolivar.livianos' y, si el texto lo indica, también a pesados/motos."))
        return cot

    # ------------------------------------------------------------------
    def _resumen(self, cot, ws, xlsx):
        rows = list(ws.iter_rows(min_row=1, max_row=60, values_only=True))
        for i, r in enumerate(rows, 1):
            t = [limpiar(c) for c in r]
            if t and t[0].startswith("Fecha de Expedición"):
                cot.fecha_cotizacion = a_fecha(next((c for c in r if a_fecha(c)), None))
            if "DATOS DEL TOMADOR" in t:
                j = t.index("DATOS DEL TOMADOR")
                cot.tomador = limpiar(rows[i][j]) if i < len(rows) else None
                cot.nit = re.sub(r"\D", "", limpiar(rows[i + 1][j])) if i + 1 < len(rows) else None
            if t and t[0] == "TIPO TASA AUTORIZADA":
                cot.tipo_tarifa = "Tasa Única" if "UNICA" in t[1].upper() else t[1].title()
            if t and t[0].startswith("Guía de Fasecolda"):
                cot.guia_fasecolda = str(r[1])
            if t and t[0] == "FORMA DE PAGO" and i < len(rows):
                cot.modalidad_pago = limpiar(rows[i][0])
            if t and t[0] == "APLICA PRIMA MÍNIMA":
                cot.notas.append(f"Aplica prima mínima: {t[1]}")
        # tabla RC / deducibles
        hdr_i = next(i for i, r in enumerate(rows) if r and limpiar(r[0]) == "CONDICIONES")
        for i in range(hdr_i + 1, hdr_i + 12):
            r = rows[i]
            cond = limpiar(r[0]).lower()
            if not cond:
                break
            key = next((k for k in _PLANES if cond.startswith(k)), None)
            if key is None:
                continue
            pid, nombre, seg = _PLANES[key]
            if not any(p.id == pid for p in cot.planes):
                cot.planes.append(Plan(id=pid, nombre=nombre, segmento=seg, descripcion=f"Condición '{limpiar(r[0])}'"))
            f = lambda col: Fuente(archivo=Path(xlsx).name, hoja=ws.title, celda=f"{col}{i + 1}")
            rc = limpiar(r[3])
            m = re.search(r"\$?\s*([\d.,]+)\s*millones", rc.lower())
            v_rc = (Valor(texto=rc, aplica=True, valor=a_numero(m.group(1)) * 1_000_000, unidad="COP", fuente=f("D")) if m
                    else Valor(texto=rc, aplica=True, valor=rc, unidad="texto", fuente=f("D")))
            cot.items.append(ItemCobertura(cobertura_id="rc_bienes", plan_id=pid, valor=v_rc))
            if "/" in rc:  # 1000/1000/2000 SMMLV
                partes = rc.replace("SMMLV", "").split("/")
                cot.items.append(ItemCobertura(cobertura_id="rc_una_persona", plan_id=pid, valor=Valor(texto=f"{partes[1].strip()} SMMLV", aplica=True, valor=a_numero(partes[1]), unidad="SMMLV", fuente=f("D"))))
                cot.items.append(ItemCobertura(cobertura_id="rc_dos_personas", plan_id=pid, valor=Valor(texto=f"{partes[2].strip()} SMMLV", aplica=True, valor=a_numero(partes[2]), unidad="SMMLV", fuente=f("D"))))
            for cid, col, idx in (("rc_bienes", "F", 5), ("ptd", "G", 6), ("pth", "G", 6), ("ppd", "H", 7), ("pph", "H", 7)):
                d = deducible(r[idx]); d.fuente = f(col)
                cot.items.append(ItemCobertura(cobertura_id=cid, plan_id=pid, tipo="deducible", valor=d))
                if cid != "rc_bienes":
                    cot.items.append(ItemCobertura(cobertura_id=cid, plan_id=pid, valor=Valor(texto="Si", aplica=True, valor=True, unidad="bool", fuente=f(col))))

    def _vehiculos(self, cot, ws, xlsx):
        rows = list(ws.iter_rows(values_only=True))
        hdr = [limpiar(c) for c in rows[0]]
        col = {n: hdr.index(n) for n in hdr if n}
        for r in rows[1:]:
            if not r or not r[col["Placa"]]:
                continue
            clase, _ = self.clase_canonica(r[col["Tipo Vehículo"]])
            cot.vehiculos.append(Vehiculo(
                placa=limpiar(r[col["Placa"]]).upper(),
                codigo_fasecolda=str(r[col["Cód. Fasecolda"]]).zfill(8) if r[col["Cód. Fasecolda"]] else None,
                marca_referencia=f"{limpiar(r[col['Marca']])} {limpiar(r[col['Referencia']])}".strip(),
                modelo=int(r[col["Modelo"]]) if r[col["Modelo"]] else None, clase=clase, ciudad=limpiar(r[col["CIUDAD"]]),
                plan=limpiar(r[col["MACRO TIPO VH"]]),
                valor_asegurado=a_numero(r[col["VrFasecolda"]]), valor_accesorios=a_numero(r[col["VrAccesorios"]]) or 0.0,
                valor_total=a_numero(r[col["Vr.Total"]]), tasa=a_tasa(r[col["Tasa"]]),
                prima_neta=a_numero(r[col["Pr neta + adicionales"]]), prima_con_iva=a_numero(r[col["Prima Total con IVA"]]),
                extra={"asistencia": a_numero(r[col["Asistencia"]]), "veh_reemplazo": a_numero(r[col["Vh Reemplazo"]]),
                       "pac": a_numero(r[col["PAC"]]), "rango_vas": limpiar(r[col["RANGO VAS"]])},
            ))

    def _tasas(self, cot, ws, xlsx):
        rows = list(ws.iter_rows(values_only=True))
        hdr = [limpiar(c) for c in rows[0]]
        for i, r in enumerate(rows[1:], start=2):
            rango = limpiar(r[0])
            if not rango:
                continue
            rango = re.sub(r"^\d+(\.\d+)?\.\s*", "", rango)  # "1. 0 - 20 Millones" -> "0 - 20 Millones"
            for j, clase in enumerate(hdr[1:], start=1):
                if clase and a_tasa(r[j]) is not None:
                    cot.tasas.append(Tasa(segmento=clase, rango=rango, tasa=a_tasa(r[j]), tipo_rango="valor_asegurado",
                                          fuente=Fuente(archivo=Path(xlsx).name, hoja=ws.title, celda=f"{chr(65 + j)}{i}")))
