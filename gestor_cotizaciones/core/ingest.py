"""
Ingesta: convierte xlsx / docx / pdf en una lista de `Bloque` (texto plano con ubicación).

Los adaptadores deterministas leen las tablas con pandas/openpyxl directamente.
Esta capa existe para la parte narrativa (slips, cláusulas, condiciones), que se
entrega al LLM como texto con referencias (hoja + fila / párrafo) para trazabilidad.

Nota: no se usan binarios externos (`pdftotext`/`pandoc`) porque no están
disponibles en la imagen Docker de este repo (ver `Dockerfile`) ni en todas las
máquinas de desarrollo. PDF se lee con `pymupdf` (ya es dependencia del repo,
mismo patrón que `conciliador.sources.foundry_recibo.extraer_texto_pdf`); el
fallback de docx con relaciones rotas lee `word/document.xml` directamente del
zip, sin pasar por `python-docx`.
"""
from __future__ import annotations

import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook

_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@dataclass
class Bloque:
    archivo: str
    seccion: str      # hoja o "docx"
    ref: str          # "fila 33" | "párrafo 12" | "tabla 3"
    texto: str

    def etiqueta(self) -> str:
        return f"[{Path(self.archivo).name} | {self.seccion} | {self.ref}]"


def bloques_xlsx(path: str | Path, hojas: list[str] | None = None, max_len: int = 400,
                  max_row: int = 500, max_col: int = 60) -> list[Bloque]:
    """`max_row`/`max_col` acotan la iteración: algunas hojas de aseguradoras (p. ej.
    "Condiciones Técnicas" del xlsx de AXA) reportan una dimensión corrupta
    (>1M filas, >16K columnas, seguramente por formato/relleno aplicado a toda la
    hoja) que sin este límite hace que `ws.iter_rows` tarde minutos por hoja."""
    wb = load_workbook(path, read_only=True, data_only=True)
    out: list[Bloque] = []
    for ws in wb.worksheets:
        if hojas and ws.title not in hojas:
            continue
        filas = ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True)
        for i, row in enumerate(filas, start=1):
            celdas = [str(c).strip() for c in row if c is not None and str(c).strip() and not str(c).startswith("#")]
            if not celdas:
                continue
            texto = " | ".join(c[:max_len] for c in celdas)
            out.append(Bloque(str(path), ws.title, f"fila {i}", texto))
    return out


def bloques_docx(path: str | Path) -> list[Bloque]:
    """Párrafos y tablas en orden de documento. python-docx primero; si el archivo tiene
    relaciones rotas (pasa con slips generados por plantillas), cae a leer
    `word/document.xml` directo del zip."""
    try:
        return _bloques_docx_python(path)
    except Exception:
        return _bloques_docx_xml(path)


def _texto_de_parrafo_xml(parrafo) -> str:
    return "".join(t.text or "" for t in parrafo.iter(f"{_W_NS}t"))


def _celdas_de_fila_xml(fila) -> list[str]:
    celdas, prev = [], None
    for celda in fila.findall(f"{_W_NS}tc"):
        texto = "\n".join(_texto_de_parrafo_xml(p) for p in celda.findall(f"{_W_NS}p")).strip()
        if texto and texto != prev:
            celdas.append(texto)
        prev = texto
    return celdas


def _bloques_docx_xml(path: str | Path) -> list[Bloque]:
    """Fallback sin binarios externos para docx con relaciones rotas (por lo que
    falla `python-docx.Document()`): lee `word/document.xml` directo del zip
    (el texto del cuerpo no depende de las relaciones rotas) y recorre
    párrafos/tablas en orden de documento, igual que `_bloques_docx_python`."""
    with zipfile.ZipFile(path) as archivo:
        xml = archivo.read("word/document.xml")
    body = ElementTree.fromstring(xml).find(f"{_W_NS}body")
    if body is None:
        return []

    out: list[Bloque] = []
    n_p = n_t = 0
    for child in body:
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            n_p += 1
            texto = _texto_de_parrafo_xml(child).strip()
            if texto:
                out.append(Bloque(str(path), "docx", f"párrafo {n_p}", texto))
        elif tag == "tbl":
            n_t += 1
            filas = [" | ".join(celdas) for fila in child.findall(f"{_W_NS}tr")
                     if (celdas := _celdas_de_fila_xml(fila))]
            if filas:
                out.append(Bloque(str(path), "docx", f"tabla {n_t}", "\n".join(filas)))
    return out


def _bloques_docx_python(path: str | Path) -> list[Bloque]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(path))
    out: list[Bloque] = []
    n_p = n_t = 0
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            n_p += 1
            t = Paragraph(child, d).text.strip()
            if t:
                out.append(Bloque(str(path), "docx", f"párrafo {n_p}", t))
        elif tag == "tbl":
            n_t += 1
            tabla = Table(child, d)
            filas = []
            for r in tabla.rows:
                celdas, prev = [], None
                for c in r.cells:
                    t = c.text.strip()
                    if t and t != prev:
                        celdas.append(t)
                    prev = t
                if celdas:
                    filas.append(" | ".join(celdas))
            if filas:
                out.append(Bloque(str(path), "docx", f"tabla {n_t}", "\n".join(filas)))
    return out


def bloques_pdf(path: str | Path) -> list[Bloque]:
    """Extrae texto por página con `pymupdf` (ya es dependencia del repo; mismo
    patrón que `conciliador.sources.foundry_recibo.extraer_texto_pdf`), en vez
    de invocar `pdftotext` (no está instalado en la imagen Docker de este repo)."""
    import pymupdf

    out: list[Bloque] = []
    with pymupdf.open(path) as documento:
        for i, pagina in enumerate(documento, start=1):
            texto = pagina.get_text().strip()
            if texto:
                out.append(Bloque(str(path), "pdf", f"página {i}", texto))
    return out


def ingerir(path: str | Path, **kw) -> list[Bloque]:
    ext = Path(path).suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        return bloques_xlsx(path, **kw)
    if ext == ".docx":
        return bloques_docx(path)
    if ext == ".pdf":
        return bloques_pdf(path)
    raise ValueError(f"Formato no soportado: {ext}")


def a_texto(bloques: list[Bloque], max_chars: int = 120_000) -> str:
    partes, total = [], 0
    for b in bloques:
        s = f"{b.etiqueta()}\n{b.texto}\n"
        if total + len(s) > max_chars:
            partes.append("[... truncado ...]")
            break
        partes.append(s)
        total += len(s)
    return "\n".join(partes)
