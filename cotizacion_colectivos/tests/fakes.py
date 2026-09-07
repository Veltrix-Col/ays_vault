from __future__ import annotations

from dataclasses import dataclass

from integrations.zoho.schemas import (
    FieldMetadata,
    ModuleMetadata,
    Organization,
    Page,
)


MODULES = (
    ModuleMetadata(
        api_name="Contacts",
        module_name="Contacts",
        plural_label="Personas",
        singular_label="Persona",
        api_supported=True,
        status="visible",
    ),
    ModuleMetadata(
        api_name="Persona_juridica",
        module_name="CustomModule8",
        plural_label="Personas juridicas",
        singular_label="Empresa",
        custom_module=True,
        api_supported=True,
        status="visible",
    ),
    ModuleMetadata(
        api_name="Polizas",
        module_name="CustomModule5",
        plural_label="Polizas",
        singular_label="Poliza",
        custom_module=True,
        api_supported=True,
        status="visible",
    ),
)

CONTACT_FIELDS = (
    FieldMetadata(
        api_name="Numero_de_documento",
        field_label="Numero de documento",
        data_type="text",
        unique=True,
    ),
    FieldMetadata(
        api_name="Full_Name",
        field_label="Nombre completo",
        data_type="text",
    ),
    FieldMetadata(
        api_name="Empresa",
        field_label="Empresa",
        data_type="lookup",
        lookup={"module": "Persona_juridica"},
    ),
)


class FakeOrganization:
    def get(self):
        return Organization(
            organization_id="sandbox-org",
            company_name="Organizacion Sandbox",
            environment="sandbox",
        )


class FakeMetadata:
    def __init__(self, failures=None):
        self.failures = failures or {}
        self.field_calls = []

    def list_modules(self):
        return MODULES

    def list_fields(self, module):
        self.field_calls.append(module)
        if module in self.failures:
            raise self.failures[module]
        if module == "Contacts":
            return CONTACT_FIELDS
        if module == "Polizas":
            return (
                FieldMetadata(
                    api_name="Tomador",
                    field_label="Tomador",
                    data_type="lookup",
                    lookup={"module": "Persona_juridica"},
                ),
            )
        return ()


class FakeRecords:
    def __init__(self):
        self.calls = []

    def list(self, **kwargs):
        self.calls.append(kwargs)
        return Page(
            records=(
                {
                    "Numero_de_documento": "123456789",
                    "Full_Name": "Nombre Real No Debe Salir",
                },
            ),
            count=1,
        )


@dataclass
class FakeZoho:
    metadata: FakeMetadata
    records: FakeRecords
    organization: FakeOrganization = FakeOrganization()
    profile: str = "sandbox"
    environment: str = "sandbox"
def mark_novelties_actions(raw):
    """Return a v4 workbook fixture with preloaded rows explicitly actionable."""
    import io
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(raw))
    for sheet in book.worksheets:
        if sheet.title in {"Póliza", "Instrucciones", "Catálogos", "Metadatos"}:
            continue
        for row in range(2, sheet.max_row + 1):
            if any(sheet.cell(row=row, column=col).value not in (None, "") for col in range(1, sheet.max_column + 1)):
                sheet.cell(row=row, column=1).value = "Ingreso"
                sheet.cell(row=row, column=2).value = "CC"
                sheet.cell(row=row, column=3).value = "123456"
                sheet.cell(row=row, column=4).value = "Persona"
                sheet.cell(row=row, column=5).value = "Prueba"
                sheet.cell(row=row, column=9).value = "2026-09-01"
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
