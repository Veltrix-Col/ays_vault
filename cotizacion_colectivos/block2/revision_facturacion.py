"""Reconstruccion READ-only del informe legacy Revision Facturacion.

Los fetchers solo emiten COQL SELECT. La reconstruccion recibe mappings o
DataFrames y hace todos los cruces localmente para conservar una regla
deterministica y comprobable sin conectarse a Zoho.
"""

from __future__ import annotations

import csv
import json
import math
from collections import Counter, defaultdict
from collections.abc import Callable, Iterable, Mapping
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from ays_zoho_sdk.exceptions import ZohoError

from cotizacion_colectivos.block2.cobros_faltantes import (
    COQL_IN_BATCH_SIZE,
    _chunks,
    _escape_coql,
    _execute_paginated,
    _fold,
    _lookup_id,
    _plain_text,
    _policy_number,
    _record_iter,
    _to_date,
    _ZOHO_ID,
    select_latest_policy_records,
)


POLICY_MODULE = "Polizas"
OPERATIONS_MODULE = "Opeeraciones"
INSURED_MODULE = "Riesgos1"
CONTACTS_MODULE = "Contacts"
NOTES_MODULE = "Notes"

POLICY_FIELDS = (
    "id", "Name", "Key", "L_nea_de_negocio", "Estado_de_la_p_liza",
    "Aseguradora1", "Ramo", "Tomador_principal1", "Analista",
    "L_der_Comercial", "Modo_de_pago", "Frecuencia",
    "Cambio_de_intermediario", "P_liza_Fecha_de_inicio_vigencia",
    "P_liza_Fecha_fin_de_la_vigencia", "Dia_facturaci_n",
    "Tipo_de_facturaci_n", "Correo_facturaci_n", "Tag",
    *(f"Fecha_{number}" for number in range(1, 13)),
)

OPERATION_FIELDS = (
    "id", "P_liza", "Name", "Observaciones",
    "Certificado_Fecha_de_inicio_de_vigencia",
    "Fecha_de_expedici_n_de_p_liza", "N_mero_de_certificado",
    "N_mero_de_certificado_Aseguradora", "Total_a_pagar_en_OP",
    "Saldo_cartera_aseguradora", "Modified_Time",
)

INSURED_FIELDS = ("id", "P_liza", "Estado", "Prima", "Pago_total", "Valor_asegurado")
CONTACT_FIELDS = ("id", "Full_Name")
NOTE_FIELDS = ("id", "Parent_Id", "Note_Content")

REPORT_COLUMNS = (
    "Key_Pol", "Aseguradora", "Ramo", "Póliza", "Estado de la póliza",
    "Etiquetas", "Nota_Op", "Nombre_Tomador", "Nombre operación",
    "Observaciones", "Periodicidad", "Dia facturación", "Tipo facturación",
    "Tipo de facturacion", "Correo de facturación", "Numero_Cuota",
    "Fecha_Cuota", "Fecha Op Inicio Vig", "Fecha Expedición",
    "Fecha_Facturacion", "Certificado (recibo, documento, anexo)",
    "Número de certificado Aseguradora", "Total a pagar en OP",
    "Saldo cartera aseguradora", "Pago total (Con IVA) Asegurados",
    "Diferencia_Cartera", "Total Valor Asegurado", "Link_Poliza", "Link_Op",
)

COMPARISON_KEY_COLUMNS = (
    "Póliza", "Nombre operación", "Numero_Cuota", "Fecha_Cuota",
    "Fecha Op Inicio Vig",
)

DATE_COLUMNS = {
    "Fecha_Cuota", "Fecha Op Inicio Vig", "Fecha Expedición", "Fecha_Facturacion",
}
NUMERIC_COLUMNS = {
    "Dia facturación", "Numero_Cuota", "Total a pagar en OP",
    "Saldo cartera aseguradora", "Pago total (Con IVA) Asegurados",
    "Diferencia_Cartera", "Total Valor Asegurado",
}

ZOHO_CRM_WEB_ORG = "753703967"
ZOHO_CRM_POLICIES_TAB = "CustomModule4"
ZOHO_CRM_OPERATIONS_TAB = "CustomModule6"

_VALID_STATES = {"vigente", "vencida"}
_VALID_BILLING_TYPES = {"mes corriente": 0, "mes vencido": 1, "mes anticipado": -1}

ReadProgress = Callable[[str], None]


class CRMReadError(Exception):
    """Agrega contexto seguro del lote sin exponer query, IDs o credenciales."""

    def __init__(
        self,
        *,
        module: str,
        filter_field: str,
        batch: int,
        total_batches: int,
        values_count: int,
        cause: ZohoError,
    ):
        super().__init__(f"Fallo READ en {module}, batch {batch}/{total_batches}.")
        self.module = module
        self.filter_field = filter_field
        self.batch = batch
        self.total_batches = total_batches
        self.values_count = values_count
        self.cause = cause


def format_read_error(exc: CRMReadError | ZohoError) -> str:
    """Renderiza solo metadata diagnostica segura disponible en ZohoError."""

    cause = exc.cause if isinstance(exc, CRMReadError) else exc
    context: list[tuple[str, object]] = []
    if isinstance(exc, CRMReadError):
        context.extend(
            (
                ("module", exc.module),
                ("filter_field", exc.filter_field),
                ("batch", exc.batch),
                ("total_batches", exc.total_batches),
                ("values", exc.values_count),
            )
        )
    elif getattr(cause, "module", ""):
        context.append(("module", cause.module))
    context.extend(
        (
            ("category", getattr(cause, "category", "unknown") or "unknown"),
            ("operation", getattr(cause, "operation", "") or "unknown"),
            ("backend", getattr(cause, "backend", "") or "unknown"),
            ("status_code", getattr(cause, "status_code", None) or "unknown"),
            ("zoho_code", getattr(cause, "zoho_code", "") or "unknown"),
            (
                "sdk_exception_class",
                getattr(cause, "sdk_exception_class", "") or "unknown",
            ),
        )
    )
    return " ".join(f"{name}={value}" for name, value in context)


@dataclass(frozen=True)
class Installment:
    number: int
    installment_date: date


@dataclass(frozen=True)
class InsuredAggregate:
    payment_total: Decimal | None
    insured_value_total: Decimal | None
    premium_total: Decimal | None
    included_records: int


@dataclass(frozen=True)
class RevisionRow:
    values: dict[str, object]
    diagnostics: tuple[str, ...] = ()


@dataclass(frozen=True)
class ReconstructionResult:
    rows: tuple[RevisionRow, ...]
    diagnostic_counts: dict[str, int]
    audit_counts: dict[str, int]
    duplicate_policy_records: int


@dataclass(frozen=True)
class ComparisonSummary:
    baseline_rows: int
    crm_rows: int
    matched_rows: int
    different_rows: int
    missing_in_crm: int
    extra_in_crm: int
    duplicate_keys: int


@dataclass(frozen=True)
class ComparisonResult:
    differences: tuple[dict[str, object], ...]
    summary: ComparisonSummary


def _materialize(records: Iterable[Mapping[str, object]] | Any) -> list[dict[str, object]]:
    return [dict(record) for record in _record_iter(records)]


def _decimal(value: object) -> Decimal | None:
    if value in (None, "") or (isinstance(value, float) and math.isnan(value)):
        return None
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        parsed = Decimal(text)
        return parsed if parsed.is_finite() else None
    except InvalidOperation:
        return None


def _output_number(value: Decimal | None) -> int | float | None:
    if value is None:
        return None
    return int(value) if value == value.to_integral_value() else float(value)


def _valid_ids(values: Iterable[object]) -> list[str]:
    supplied = {str(value).strip() for value in values if str(value or "").strip()}
    if any(not _ZOHO_ID.fullmatch(value) for value in supplied):
        raise ValueError("Zoho devolvio uno o mas IDs invalidos.")
    return sorted(supplied)


def _fetch_by_in(
    zoho: Any,
    *,
    module: str,
    fields: tuple[str, ...],
    filter_field: str,
    values: Iterable[str],
    values_are_ids: bool = True,
    progress: ReadProgress | None = None,
) -> list[dict[str, object]]:
    unique = _valid_ids(values) if values_are_ids else sorted({str(value).strip() for value in values if str(value).strip()})
    records: list[dict[str, object]] = []
    selected = ", ".join(fields)
    batches = list(_chunks(unique, COQL_IN_BATCH_SIZE))
    for batch_number, batch in enumerate(batches, start=1):
        if progress:
            progress(
                f"[READ] {module}: batch {batch_number}/{len(batches)}, "
                f"filter={filter_field}, values={len(batch)}"
            )
        rendered = ", ".join(f"'{_escape_coql(value)}'" for value in batch)
        query = f"select {selected} from {module} where {filter_field} in ({rendered})"
        try:
            batch_records = _execute_paginated(zoho, query)
        except ZohoError as exc:
            raise CRMReadError(
                module=module,
                filter_field=filter_field,
                batch=batch_number,
                total_batches=len(batches),
                values_count=len(batch),
                cause=exc,
            ) from exc
        records.extend(batch_records)
        if progress:
            progress(
                f"[READ] {module}: batch {batch_number}/{len(batches)} OK, "
                f"{len(batch_records)} registros"
            )
    return records


def fetch_policy_records(
    zoho: Any,
    policy_numbers: Iterable[str],
    *,
    progress: ReadProgress | None = None,
) -> list[dict[str, object]]:
    return _fetch_by_in(
        zoho,
        module=POLICY_MODULE,
        fields=POLICY_FIELDS,
        filter_field="Name",
        values=(_policy_number(value) for value in policy_numbers),
        values_are_ids=False,
        progress=progress,
    )


def fetch_operation_records(
    zoho: Any,
    policy_ids: Iterable[str],
    *,
    progress: ReadProgress | None = None,
) -> list[dict[str, object]]:
    return _fetch_by_in(
        zoho, module=OPERATIONS_MODULE, fields=OPERATION_FIELDS,
        filter_field="P_liza", values=policy_ids, progress=progress,
    )


def fetch_insured_records(
    zoho: Any,
    policy_ids: Iterable[str],
    *,
    progress: ReadProgress | None = None,
) -> list[dict[str, object]]:
    return _fetch_by_in(
        zoho, module=INSURED_MODULE, fields=INSURED_FIELDS,
        filter_field="P_liza", values=policy_ids, progress=progress,
    )


def fetch_contact_records(
    zoho: Any,
    contact_ids: Iterable[str],
    *,
    progress: ReadProgress | None = None,
) -> list[dict[str, object]]:
    return _fetch_by_in(
        zoho, module=CONTACTS_MODULE, fields=CONTACT_FIELDS,
        filter_field="id", values=contact_ids, progress=progress,
    )


def fetch_note_records(
    zoho: Any,
    operation_ids: Iterable[str],
    *,
    progress: ReadProgress | None = None,
) -> list[dict[str, object]]:
    return _fetch_by_in(
        zoho, module=NOTES_MODULE, fields=NOTE_FIELDS,
        filter_field="Parent_Id", values=operation_ids, progress=progress,
    )


def is_legacy_policy_candidate(policy: Mapping[str, object]) -> bool:
    return (
        _fold(policy.get("L_nea_de_negocio")) == "colectivos"
        and _fold(policy.get("Estado_de_la_p_liza")) in _VALID_STATES
        and _fold(policy.get("Modo_de_pago")) == "fraccionado"
    )


def select_latest_candidate_policies(
    policies: Iterable[Mapping[str, object]] | Any,
) -> tuple[dict[str, dict[str, object]], dict[str, int]]:
    materialized = _materialize(policies)
    _all_selected, duplicates = select_latest_policy_records(materialized)
    eligible = [policy for policy in materialized if is_legacy_policy_candidate(policy)]
    selected, _eligible_duplicates = select_latest_policy_records(eligible)
    return selected, duplicates


def unpivot_payment_plan(policy: Mapping[str, object]) -> tuple[Installment, ...]:
    installments = []
    for number in range(1, 13):
        installment_date = _to_date(policy.get(f"Fecha_{number}"))
        if installment_date is not None:
            installments.append(Installment(number, installment_date))
    return tuple(installments)


def _shift_month(year: int, month: int, shift: int) -> tuple[int, int]:
    absolute = year * 12 + (month - 1) + shift
    return absolute // 12, absolute % 12 + 1


def calculate_billing_date(
    base_date: object,
    billing_day: object,
    billing_type: object,
) -> date | None:
    base = _to_date(base_date)
    day_decimal = _decimal(billing_day)
    folded_type = _fold(billing_type)
    if base is None or day_decimal is None or day_decimal != day_decimal.to_integral_value():
        return None
    day = int(day_decimal)
    if day < 1 or day > 31 or folded_type not in _VALID_BILLING_TYPES:
        return None
    shift = _VALID_BILLING_TYPES[folded_type]
    if base.day > day:
        shift += 1
    year, month = _shift_month(base.year, base.month, shift)
    try:
        return date(year, month, day)
    except ValueError:
        # No se recorta al ultimo dia: el SQL legacy no esta disponible para
        # demostrar ese comportamiento.
        return None


def _billing_diagnostics(policy: Mapping[str, object], base_date: date) -> set[str]:
    diagnostics: set[str] = set()
    billing_day = policy.get("Dia_facturaci_n")
    billing_type = policy.get("Tipo_de_facturaci_n")
    if billing_day in (None, ""):
        diagnostics.add("MISSING_BILLING_DAY")
    if not _plain_text(billing_type):
        diagnostics.add("MISSING_BILLING_TYPE")
    if calculate_billing_date(base_date, billing_day, billing_type) is None and not diagnostics:
        diagnostics.add("INVALID_BILLING_DAY")
    return diagnostics


def is_prorroga(operation: Mapping[str, object]) -> bool:
    return (
        _fold(operation.get("Name")).startswith("modifi")
        and "prorroga" in _fold(operation.get("Observaciones"))
    )


def is_pending_operation(operation: Mapping[str, object] | None) -> bool:
    if operation is None:
        return True
    certificate = operation.get("N_mero_de_certificado")
    expedition = operation.get("Fecha_de_expedici_n_de_p_liza")
    total = _decimal(operation.get("Total_a_pagar_en_OP"))
    certificate_missing = certificate is None or not str(certificate).strip()
    expedition_missing = expedition is None or not str(expedition).strip()
    return certificate_missing or expedition_missing or total == Decimal("0")


def is_in_accumulated_current_month(value: object, *, as_of: date) -> bool:
    candidate = _to_date(value)
    return candidate is not None and candidate.year == as_of.year and candidate.month <= as_of.month


def select_relevant_operation_ids_for_notes(
    policies: Iterable[Mapping[str, object]] | Any,
    operations: Iterable[Mapping[str, object]] | Any,
    *,
    as_of: date,
) -> tuple[str, ...]:
    """Reduce Notes a operaciones que pueden producir una fila final.

    Notes no participa en los filtros de entrada, por lo que consultar notas
    de operaciones completas, fuera del corte o no asociadas al plan solo
    aumenta volumen sin cambiar el resultado reconstruido.
    """

    selected, _duplicates = select_latest_candidate_policies(policies)
    operations_by_policy: dict[str, list[dict[str, object]]] = defaultdict(list)
    for operation in _materialize(operations):
        operations_by_policy[_lookup_id(operation.get("P_liza"))].append(operation)

    relevant: set[str] = set()

    def include_if_reportable(
        policy: Mapping[str, object],
        operation: Mapping[str, object],
        base_date: object,
    ) -> None:
        billing_date = calculate_billing_date(
            base_date,
            policy.get("Dia_facturaci_n"),
            policy.get("Tipo_de_facturaci_n"),
        )
        operation_id = str(operation.get("id") or "").strip()
        if (
            operation_id
            and is_pending_operation(operation)
            and is_in_accumulated_current_month(billing_date, as_of=as_of)
        ):
            relevant.add(operation_id)

    for policy in selected.values():
        policy_operations = operations_by_policy.get(str(policy.get("id") or ""), [])
        for installment in unpivot_payment_plan(policy):
            for operation in policy_operations:
                if (
                    not _fold(operation.get("Name")).startswith("modifi")
                    and _to_date(
                        operation.get("Certificado_Fecha_de_inicio_de_vigencia")
                    )
                    == installment.installment_date
                ):
                    include_if_reportable(
                        policy, operation, installment.installment_date
                    )
        for operation in policy_operations:
            if is_prorroga(operation):
                include_if_reportable(
                    policy,
                    operation,
                    operation.get("Certificado_Fecha_de_inicio_de_vigencia"),
                )

    return tuple(sorted(relevant))


def aggregate_insured_by_policy(
    records: Iterable[Mapping[str, object]] | Any,
) -> dict[str, InsuredAggregate]:
    grouped: dict[str, list[dict[str, object]]] = defaultdict(list)
    for record in _materialize(records):
        state = _fold(record.get("Estado"))
        if "activo" in state or state == "excluido con cobro":
            grouped[_lookup_id(record.get("P_liza"))].append(record)

    result: dict[str, InsuredAggregate] = {}
    for policy_id, included in grouped.items():
        def summed(field: str) -> Decimal | None:
            values = [_decimal(item.get(field)) for item in included]
            numeric = [value for value in values if value is not None]
            return sum(numeric, Decimal("0")) if numeric else None

        result[policy_id] = InsuredAggregate(
            payment_total=summed("Pago_total"),
            insured_value_total=summed("Valor_asegurado"),
            premium_total=summed("Prima"),
            included_records=len(included),
        )
    return result


def policy_url(policy_id: object) -> str:
    value = str(policy_id or "").strip()
    return f"https://crm.zoho.com/crm/org{ZOHO_CRM_WEB_ORG}/tab/{ZOHO_CRM_POLICIES_TAB}/{value}" if _ZOHO_ID.fullmatch(value) else ""


def operation_url(operation_id: object) -> str:
    value = str(operation_id or "").strip()
    return f"https://crm.zoho.com/crm/org{ZOHO_CRM_WEB_ORG}/tab/{ZOHO_CRM_OPERATIONS_TAB}/{value}" if _ZOHO_ID.fullmatch(value) else ""


def _tags(value: object) -> str:
    if isinstance(value, (list, tuple)):
        names = [_plain_text(item) for item in value]
        return ",".join(name for name in names if name)
    return _plain_text(value)


def _row_values(
    policy: Mapping[str, object],
    operation: Mapping[str, object] | None,
    *,
    note: Mapping[str, object] | None,
    contact_names: Mapping[str, str],
    insured: InsuredAggregate | None,
    installment: Installment | None,
    billing_date: date | None,
) -> dict[str, object]:
    operation = operation or {}
    payment_total = insured.payment_total if insured else None
    insured_value = insured.insured_value_total if insured else None
    balance = _decimal(operation.get("Saldo_cartera_aseguradora"))
    # El baseline real contiene tres agregados de asegurados vacios y en los
    # tres Diferencia_Cartera=0 cuando el saldo tambien esta vacio. Esto
    # demuestra COALESCE a cero para ambos lados dentro de esta expresion.
    difference = abs(
        (payment_total or Decimal("0")) - (balance or Decimal("0"))
    )
    policy_id = str(policy.get("id") or "")
    operation_id = str(operation.get("id") or "")
    billing_type = _plain_text(policy.get("Tipo_de_facturaci_n"))
    return {
        "Key_Pol": _plain_text(policy.get("Key")),
        "Aseguradora": _plain_text(policy.get("Aseguradora1")),
        "Ramo": _plain_text(policy.get("Ramo")),
        "Póliza": _policy_number(policy.get("Name")),
        "Estado de la póliza": _plain_text(policy.get("Estado_de_la_p_liza")),
        "Etiquetas": _tags(policy.get("Tag")),
        "Nota_Op": _plain_text((note or {}).get("Note_Content")),
        "Nombre_Tomador": contact_names.get(_lookup_id(policy.get("Tomador_principal1")), ""),
        "Nombre operación": _plain_text(operation.get("Name")),
        "Observaciones": _plain_text(operation.get("Observaciones")),
        "Periodicidad": _plain_text(policy.get("Frecuencia")),
        "Dia facturación": _output_number(_decimal(policy.get("Dia_facturaci_n"))),
        # No existe SQL local que demuestre dos fuentes distintas. Se conservan
        # ambas columnas y se alimentan explicitamente desde el campo confirmado.
        "Tipo facturación": billing_type,
        "Tipo de facturacion": billing_type,
        "Correo de facturación": _plain_text(policy.get("Correo_facturaci_n")),
        "Numero_Cuota": installment.number if installment else None,
        "Fecha_Cuota": installment.installment_date if installment else None,
        "Fecha Op Inicio Vig": _to_date(operation.get("Certificado_Fecha_de_inicio_de_vigencia")),
        "Fecha Expedición": _to_date(operation.get("Fecha_de_expedici_n_de_p_liza")),
        "Fecha_Facturacion": billing_date,
        "Certificado (recibo, documento, anexo)": _plain_text(operation.get("N_mero_de_certificado")),
        "Número de certificado Aseguradora": _plain_text(operation.get("N_mero_de_certificado_Aseguradora")),
        "Total a pagar en OP": _output_number(_decimal(operation.get("Total_a_pagar_en_OP"))),
        "Saldo cartera aseguradora": _output_number(balance),
        "Pago total (Con IVA) Asegurados": _output_number(payment_total),
        "Diferencia_Cartera": _output_number(difference),
        "Total Valor Asegurado": _output_number(insured_value),
        "Link_Poliza": policy_url(policy_id),
        "Link_Op": operation_url(operation_id),
    }


def reconstruct_revision_facturacion(
    policies: Iterable[Mapping[str, object]] | Any,
    operations: Iterable[Mapping[str, object]] | Any,
    insured_records: Iterable[Mapping[str, object]] | Any,
    contacts: Iterable[Mapping[str, object]] | Any,
    notes: Iterable[Mapping[str, object]] | Any,
    *,
    as_of: date,
) -> ReconstructionResult:
    selected, duplicates = select_latest_candidate_policies(policies)
    operation_list = _materialize(operations)
    operations_by_policy: dict[str, list[dict[str, object]]] = defaultdict(list)
    for operation in operation_list:
        operations_by_policy[_lookup_id(operation.get("P_liza"))].append(operation)
    notes_by_operation: dict[str, list[dict[str, object]]] = defaultdict(list)
    for note in _materialize(notes):
        notes_by_operation[_lookup_id(note.get("Parent_Id"))].append(note)
    contact_names = {
        str(contact.get("id") or ""): _plain_text(contact.get("Full_Name"))
        for contact in _materialize(contacts)
    }
    insured_by_policy = aggregate_insured_by_policy(insured_records)

    rows: list[RevisionRow] = []
    diagnostic_counts: Counter[str] = Counter()
    audit_counts: Counter[str] = Counter()

    def emit(
        policy: Mapping[str, object],
        operation: Mapping[str, object] | None,
        installment: Installment | None,
        base_diagnostics: set[str],
    ) -> None:
        base_date = installment.installment_date if installment else _to_date((operation or {}).get("Certificado_Fecha_de_inicio_de_vigencia"))
        diagnostics = set(base_diagnostics)
        if base_date is None:
            diagnostics.add("INVALID_BILLING_DAY")
            for code in diagnostics:
                diagnostic_counts[code] += 1
            return
        diagnostics.update(_billing_diagnostics(policy, base_date))
        billing_date = calculate_billing_date(
            base_date, policy.get("Dia_facturaci_n"), policy.get("Tipo_de_facturaci_n")
        )
        if billing_date and billing_date > as_of and billing_date.year == as_of.year and billing_date.month == as_of.month:
            diagnostics.add("FUTURE_IN_CURRENT_MONTH")
        policy_id = str(policy.get("id") or "")
        insured = insured_by_policy.get(policy_id)
        if insured is None or insured.payment_total is None:
            diagnostics.add("NO_INSURED_TOTAL")
        operation_id = str((operation or {}).get("id") or "")
        related_notes = notes_by_operation.get(operation_id, []) if operation_id else []
        if len(related_notes) > 1:
            diagnostics.add("MULTIPLE_NOTES")
        note_variants: list[Mapping[str, object] | None] = related_notes or [None]
        for code in diagnostics:
            diagnostic_counts[code] += 1
        if installment is not None and installment.installment_date.year != as_of.year:
            return
        if not is_pending_operation(operation) or not is_in_accumulated_current_month(billing_date, as_of=as_of):
            return
        if "MISSING_OPERATION" in diagnostics:
            audit_counts["MISSING_OPERATION_TEMPORAL_SURVIVORS"] += 1
        if "PRORROGA" in diagnostics:
            audit_counts["PRORROGA_TEMPORAL_SURVIVORS"] += 1
        for note in note_variants:
            values = _row_values(
                policy, operation, note=note, contact_names=contact_names,
                insured=insured, installment=installment, billing_date=billing_date,
            )
            rows.append(RevisionRow(values, tuple(sorted(diagnostics))))
            if "MISSING_OPERATION" in diagnostics:
                audit_counts["MISSING_OPERATION_FINAL_ROWS"] += 1
            if "PRORROGA" in diagnostics:
                audit_counts["PRORROGA_FINAL_ROWS"] += 1

    for name, policy in selected.items():
        policy_id = str(policy.get("id") or "")
        policy_operations = operations_by_policy.get(policy_id, [])
        shared = {"MULTIPLE_POLICY_RECORDS"} if name in duplicates else set()
        start = _to_date(policy.get("P_liza_Fecha_de_inicio_vigencia"))
        end = _to_date(policy.get("P_liza_Fecha_fin_de_la_vigencia"))
        for installment in unpivot_payment_plan(policy):
            audit_counts["PLAN_INSTALLMENT_CANDIDATES"] += 1
            diagnostics = set(shared)
            if (start and installment.installment_date < start) or (end and installment.installment_date > end):
                diagnostics.add("PLAN_DATE_OUTSIDE_POLICY_TERM")
            matches = [
                operation for operation in policy_operations
                if not _fold(operation.get("Name")).startswith("modifi")
                and _to_date(operation.get("Certificado_Fecha_de_inicio_de_vigencia")) == installment.installment_date
            ]
            if not matches:
                audit_counts["MISSING_OPERATION_CANDIDATES"] += 1
                diagnostics.add("MISSING_OPERATION")
                emit(policy, None, installment, diagnostics)
            else:
                if len(matches) > 1:
                    diagnostics.add("MULTIPLE_OPERATIONS_FOR_INSTALLMENT")
                for operation in matches:
                    emit(policy, operation, installment, diagnostics)

        for operation in policy_operations:
            if is_prorroga(operation):
                audit_counts["PRORROGA_CANDIDATES"] += 1
                emit(policy, operation, None, {*shared, "PRORROGA"})

    return ReconstructionResult(
        rows=tuple(rows),
        diagnostic_counts=dict(sorted(diagnostic_counts.items())),
        audit_counts=dict(sorted(audit_counts.items())),
        duplicate_policy_records=len(duplicates),
    )


def _header_key(value: object) -> str:
    return "".join(_fold(value).split())


def read_baseline(path: str | Path) -> list[dict[str, object]]:
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)
        try:
            raw_headers = next(iterator)
        except StopIteration as exc:
            raise ValueError("El baseline no contiene encabezados.") from exc
        indexes = {_header_key(value): index for index, value in enumerate(raw_headers)}
        missing = [column for column in REPORT_COLUMNS if _header_key(column) not in indexes]
        if missing:
            raise ValueError(f"Faltan columnas requeridas: {', '.join(missing)}.")
        rows = []
        for raw in iterator:
            row = {column: raw[indexes[_header_key(column)]] for column in REPORT_COLUMNS}
            if all(value in (None, "") for value in row.values()):
                continue
            row["Póliza"] = _policy_number(row["Póliza"])
            rows.append(row)
        return rows
    finally:
        workbook.close()


def _normalized_value(column: str, value: object) -> object:
    if (
        value is None
        or (isinstance(value, float) and math.isnan(value))
        or (isinstance(value, str) and not value.strip())
    ):
        return None
    if column in DATE_COLUMNS:
        parsed = _to_date(value)
        return parsed.isoformat() if parsed else _plain_text(value)
    if column in NUMERIC_COLUMNS:
        number = _decimal(value)
        if number is None:
            return None
        if number == number.to_integral_value():
            return number.quantize(Decimal("1"))
        return number.normalize()
    if column == "Póliza":
        return _policy_number(value)
    if column == "Etiquetas":
        # Zoho entrega una secuencia; Analytics la serializa separada por
        # comas. Se ignora solo el espacio alrededor del separador y se
        # conserva el orden para no ocultar cambios semánticos.
        return ",".join(
            part.strip() for part in _plain_text(value).split(",") if part.strip()
        )
    return _plain_text(value)


def _comparison_key(row: Mapping[str, object]) -> tuple[object, ...]:
    return tuple(_normalized_value(column, row.get(column)) for column in COMPARISON_KEY_COLUMNS)


def _signature(row: Mapping[str, object]) -> tuple[object, ...]:
    return tuple(_normalized_value(column, row.get(column)) for column in REPORT_COLUMNS)


def _json_value(value: object) -> object:
    if isinstance(value, (date, Decimal)):
        return str(value)
    return value


def compare_with_baseline(
    baseline_rows: Iterable[Mapping[str, object]],
    crm_rows: Iterable[RevisionRow],
) -> ComparisonResult:
    baseline = [dict(row) for row in baseline_rows]
    crm = list(crm_rows)
    baseline_by_key: dict[tuple[object, ...], list[dict[str, object]]] = defaultdict(list)
    crm_by_key: dict[tuple[object, ...], list[RevisionRow]] = defaultdict(list)
    for row in baseline:
        baseline_by_key[_comparison_key(row)].append(row)
    for row in crm:
        crm_by_key[_comparison_key(row.values)].append(row)

    duplicate_keys = sum(
        1 for key in set(baseline_by_key) | set(crm_by_key)
        if len(baseline_by_key.get(key, ())) > 1 or len(crm_by_key.get(key, ())) > 1
    )
    matched = different = missing = extra = 0
    differences: list[dict[str, object]] = []

    def detail(
        status: str,
        key: tuple[object, ...],
        baseline_row: Mapping[str, object] | None,
        crm_row: RevisionRow | None,
    ) -> dict[str, object]:
        baseline_row = baseline_row or {}
        crm_values = crm_row.values if crm_row else {}
        changed = [
            column for column in REPORT_COLUMNS
            if _normalized_value(column, baseline_row.get(column)) != _normalized_value(column, crm_values.get(column))
        ]
        result: dict[str, object] = {
            "STATUS": status,
            "KEY": json.dumps([_json_value(value) for value in key], ensure_ascii=False),
            "DIFFERING_COLUMNS": ";".join(changed),
            "CRM_DIAGNOSTICS": ";".join(crm_row.diagnostics) if crm_row else "",
        }
        for column in REPORT_COLUMNS:
            result[f"BASELINE_{column}"] = baseline_row.get(column)
            result[f"CRM_{column}"] = crm_values.get(column)
        return result

    for key in sorted(set(baseline_by_key) | set(crm_by_key), key=str):
        base_group = sorted(baseline_by_key.get(key, []), key=lambda row: str(_signature(row)))
        crm_group = sorted(crm_by_key.get(key, []), key=lambda row: str(_signature(row.values)))
        unmatched_crm = list(crm_group)
        unmatched_base = []
        for base_row in base_group:
            signature = _signature(base_row)
            exact_index = next(
                (index for index, crm_row in enumerate(unmatched_crm) if _signature(crm_row.values) == signature),
                None,
            )
            if exact_index is None:
                unmatched_base.append(base_row)
            else:
                matched += 1
                unmatched_crm.pop(exact_index)
        pairs = min(len(unmatched_base), len(unmatched_crm))
        for index in range(pairs):
            different += 1
            differences.append(detail("DIFFERENT", key, unmatched_base[index], unmatched_crm[index]))
        for base_row in unmatched_base[pairs:]:
            missing += 1
            differences.append(detail("MISSING_IN_CRM", key, base_row, None))
        for crm_row in unmatched_crm[pairs:]:
            extra += 1
            differences.append(detail("EXTRA_IN_CRM", key, None, crm_row))

    return ComparisonResult(
        differences=tuple(differences),
        summary=ComparisonSummary(
            baseline_rows=len(baseline), crm_rows=len(crm), matched_rows=matched,
            different_rows=different, missing_in_crm=missing, extra_in_crm=extra,
            duplicate_keys=duplicate_keys,
        ),
    )


def write_differences_csv(path: str | Path, differences: Iterable[Mapping[str, object]]) -> None:
    fields = (
        "STATUS", "KEY", "DIFFERING_COLUMNS", "CRM_DIAGNOSTICS",
        *(f"BASELINE_{column}" for column in REPORT_COLUMNS),
        *(f"CRM_{column}" for column in REPORT_COLUMNS),
    )
    with Path(path).open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(differences)

