from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook

from .services import _current_and_historical_text


MAX_CELL = 32000
HUMAN_COLUMNS = ("VALIDACION_HUMANA", "DECISION_HUMANA", "COMENTARIO_HUMANO")


def _value(value):
    return "" if value is None else value


def _text(value, limit=MAX_CELL):
    value = "" if value is None else str(value)
    return value if len(value) <= limit else value[:limit] + "…"


def _bool(value):
    return str(value or "").strip().lower() in {"1", "true", "yes", "si", "sí"}


def _load_sheet(path, sheet_name):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "") for value in next(rows)]
        return [dict(zip(headers, row)) for row in rows]
    finally:
        workbook.close()


def _load_replay(path):
    return _load_sheet(path, "ComparacionV631")


def _load_reference(path):
    rows = _load_sheet(path, "Cambios")
    return {str(row.get("message_id_hash") or ""): row for row in rows if row.get("message_id_hash")}


def _terms_for_rule(rule):
    return {
        "SIGNATURE_FAILURE_V631": ("firma", "rechazad", "error"),
        "PAYMENT_LEGALIZATION_V631": ("legalizacion", "legalizar", "pago"),
        "COLLECTION_DOCUMENTATION_V631": ("cobro", "document", "pendiente", "falta", "solicitud"),
        "EMISSION_INCONSISTENCY_V631": ("emision", "inconsist", "operacion poliza nueva"),
        "RESCHEDULE_REQUIRED_V631": ("reagendar", "reagendamiento", "nueva cita"),
        "PAYMENT_RECEIPT_CONTEXTUAL_V631": ("comprobante de pago", "comprobante pago"),
        "INVOICE_MISSING_V631": ("factura", "recibido", "pendiente", "solicitud", "falta"),
        "PAYMENT_DUPLICATE_V631": ("pago duplicado", "pago doble", "doble pago", "cobro duplicado"),
    }.get(rule, ())


def _text_context(row):
    subject = row.get("subject") or ""
    body = row.get("body_text") or ""
    current, historical = _current_and_historical_text(subject, body)
    return current, historical


def _signals(row, rule):
    terms = _terms_for_rule(rule)
    current, historical = _text_context(row)
    return (
        tuple(term for term in terms if term in current),
        tuple(term for term in terms if term in historical and term not in current),
    )


def _root_cause(row, reference):
    current_rule = str(row.get("RULE_ID") or "")
    ref_outcome = str(reference.get("V631_MESSAGE_OUTCOME") or "")
    ref_scope = str(reference.get("V631_SCOPE") or "")
    current_scope = str(row.get("SCOPE") or "")
    current_queue = _bool(row.get("CURRENT_QUEUE"))
    reference_queue = _bool(row.get("V631_QUEUE"))
    if current_scope != ref_scope and ref_scope:
        return "SCOPE_DIFFERENCE", "RULE_SCOPE_DIFFERENCE"
    if reference.get("V631_CASE_STATUS") in {"RESOLVED", "CLOSED", "RESUELTO", "CERRADO"} and current_queue:
        return "CASE_STATE", "REFERENCE_CASE_RESOLVED"
    current_signals, historical_signals = _signals(row, current_rule)
    if historical_signals and not current_signals:
        return "HISTORICAL_TEXT", "RULE_SIGNAL_ONLY_IN_QUOTED_HISTORY"
    if not reference_queue and current_queue:
        if current_rule in {"SIGNATURE_FAILURE_V631", "PAYMENT_LEGALIZATION_V631", "COLLECTION_DOCUMENTATION_V631", "EMISSION_INCONSISTENCY_V631", "RESCHEDULE_REQUIRED_V631"}:
            return "RULE_TOO_BROAD", "CURRENT_RULE_MATCHES_NON_QUEUE_REFERENCE"
        return "REFERENCE_V631_QUESTIONABLE", "CURRENT_SEMANTIC_MATCH_REFERENCE_NOT_QUEUE"
    if reference_queue and not current_queue:
        if current_rule == "UNCLEAR_V631":
            return "VOCABULARY_GAP", "REFERENCE_ACTION_NOT_RECOGNIZED"
        if current_rule in {"INFORMATIONAL_V631", "COMMUNICATIONS_NOISE_V631"}:
            return "RULE_TOO_NARROW", "CURRENT_NEGATIVE_CLASSIFICATION"
        return "AMBIGUOUS", "REFERENCE_QUEUE_CURRENT_NOT_ELIGIBLE"
    if ref_outcome == "UNCLEAR" or row.get("MESSAGE_OUTCOME") == "UNCLEAR":
        return "AMBIGUOUS", "UNCLEAR_OUTCOME"
    return "OTHER", "UNRESOLVED_DIFFERENCE"


def _recommendation(root_cause, row):
    if root_cause == "RULE_TOO_BROAD":
        return "TIGHTEN_CURRENT_RULE", "MEDIUM"
    if root_cause in {"VOCABULARY_GAP", "RULE_TOO_NARROW"}:
        return "EXPAND_CURRENT_RULE", "LOW"
    if root_cause in {"CASE_STATE", "HISTORICAL_TEXT"}:
        return "MOVE_TO_CASE_LAYER" if root_cause == "CASE_STATE" else "KEEP_CURRENT", "MEDIUM"
    if root_cause in {"REFERENCE_V631_QUESTIONABLE", "AMBIGUOUS"}:
        return "NEEDS_HUMAN_REVIEW", "LOW"
    if row.get("CURRENT_QUEUE") and row.get("V631_QUEUE"):
        return "KEEP_CURRENT", "HIGH"
    return "NEEDS_HUMAN_REVIEW", "LOW"


def _diagnostic_row(row, reference):
    root, detail = _root_cause(row, reference)
    recommendation, confidence = _recommendation(root, row)
    current_signals, historical_signals = _signals(row, row.get("RULE_ID") or "")
    if row.get("QUEUE_RESULT") == "POTENTIAL_FALSE_POSITIVE":
        divergence_type = "FALSE_POSITIVE"
    elif row.get("QUEUE_RESULT") == "POTENTIAL_FALSE_NEGATIVE":
        divergence_type = "FALSE_NEGATIVE"
    else:
        divergence_type = "OTHER_DIVERGENCE"
    return {
        "DIVERGENCE_TYPE": divergence_type,
        "message_id_hash": row.get("message_id_hash", ""),
        "received_at": row.get("received_at", ""),
        "from_email": row.get("sender", ""),
        "subject": row.get("subject", ""),
        "body_preview": _text(row.get("body_text", ""), 1000),
        "CURRENT_RULE_ID": row.get("RULE_ID", ""),
        "CURRENT_FAMILY": row.get("FAMILY", ""),
        "CURRENT_EVENT_TYPE": row.get("EVENT_TYPE", ""),
        "CURRENT_ACTION_TYPE": row.get("ACTION_TYPE", ""),
        "CURRENT_SCOPE": row.get("SCOPE", ""),
        "CURRENT_MESSAGE_OUTCOME": row.get("MESSAGE_OUTCOME", ""),
        "CURRENT_ACTION_REQUIRED": row.get("ACTION_REQUIRED", ""),
        "CURRENT_QUEUE": row.get("CURRENT_QUEUE", ""),
        "CURRENT_ORGANIZATION": row.get("ORGANIZATION", ""),
        "REFERENCE_RULE_ID": reference.get("V631_RULE_ID", ""),
        "REFERENCE_FAMILY": reference.get("V631_FAMILY", ""),
        "REFERENCE_EVENT": reference.get("V631_EVENT", ""),
        "REFERENCE_ACTION_TYPE": reference.get("V631_ACTION_TYPE", ""),
        "REFERENCE_SCOPE": reference.get("V631_SCOPE", ""),
        "REFERENCE_MESSAGE_OUTCOME": reference.get("V631_MESSAGE_OUTCOME", ""),
        "REFERENCE_ACTION_REQUIRED": reference.get("V631_ACTION_REQUIRED", ""),
        "REFERENCE_QUEUE": reference.get("V631_SHOULD_BE_IN_EXCEPTION_QUEUE", ""),
        "REFERENCE_CASE_STATUS": reference.get("V631_CASE_STATUS", ""),
        "REFERENCE_ORGANIZATION": reference.get("V631_ORGANIZATION", ""),
        "CURRENT_SIGNALS": ", ".join(current_signals),
        "HISTORICAL_ONLY_SIGNALS": ", ".join(historical_signals),
        "ROOT_CAUSE": root,
        "SECONDARY_CAUSE": detail,
        "RECOMMENDATION": recommendation,
        "CONFIDENCE_OF_RECOMMENDATION": confidence,
    }


def _group_rows(rows, fields):
    groups = defaultdict(lambda: {"COUNT": 0})
    for row in rows:
        key = tuple(row.get(field, "") for field in fields)
        group = groups[key]
        group["COUNT"] += 1
        for field, value in zip(fields, key):
            group[field] = value
    return list(groups.values())


def _write_sheet(workbook, title, rows):
    sheet = workbook.create_sheet(title)
    if not rows:
        sheet.append(["No data"])
        return
    headers = list(rows[0])
    sheet.append(headers)
    for row in rows:
        sheet.append([_text(row.get(header)) for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _sample(rows, limit=150):
    priority = {
        "SIGNATURE_FAILURE_V631": 0,
        "PAYMENT_LEGALIZATION_V631": 1,
        "FALSE_NEGATIVE": 2,
        "HISTORICAL_TEXT": 3,
        "CASE_STATE": 4,
        "AMBIGUOUS": 5,
    }
    ranked = sorted(rows, key=lambda row: (priority.get(row.get("CURRENT_RULE_ID"), 6), row.get("ROOT_CAUSE", ""), row.get("message_id_hash", "")))
    selected = ranked[:limit]
    for row in selected:
        for column in HUMAN_COLUMNS:
            row[column] = ""
    return selected


def run_diagnostic(replay_path, reference_path, output_path):
    replay_rows = _load_replay(replay_path)
    references = _load_reference(reference_path)
    diagnostic_rows = [_diagnostic_row(row, references.get(str(row.get("message_id_hash") or ""), {})) for row in replay_rows if row.get("QUEUE_RESULT") in {"POTENTIAL_FALSE_POSITIVE", "POTENTIAL_FALSE_NEGATIVE"}]
    false_positives = [row for row in diagnostic_rows if row["DIVERGENCE_TYPE"] == "FALSE_POSITIVE"]
    false_negatives = [row for row in diagnostic_rows if row["DIVERGENCE_TYPE"] == "FALSE_NEGATIVE"]
    unclear = [
        _diagnostic_row(row, references.get(str(row.get("message_id_hash") or ""), {}))
        for row in replay_rows
        if row.get("MESSAGE_OUTCOME") == "UNCLEAR"
    ]
    for row in unclear:
        row["UNCLEAR_REFERENCE_BUCKET"] = row.get("REFERENCE_MESSAGE_OUTCOME") or "NO_REFERENCE_OUTCOME"
    signature = [row for row in false_positives if row["CURRENT_RULE_ID"] == "SIGNATURE_FAILURE_V631"]
    payment = [row for row in false_positives if row["CURRENT_RULE_ID"] == "PAYMENT_LEGALIZATION_V631"]
    temporal = [row for row in diagnostic_rows if row["ROOT_CAUSE"] == "HISTORICAL_TEXT"]
    case_state = [row for row in diagnostic_rows if row["ROOT_CAUSE"] == "CASE_STATE"]
    organization = [row for row in diagnostic_rows if row["ROOT_CAUSE"] == "ORGANIZATION_DIFFERENCE"]
    root_counts = Counter(row["ROOT_CAUSE"] for row in diagnostic_rows)
    recommendations = Counter((row["CURRENT_RULE_ID"], row["RECOMMENDATION"], row["CONFIDENCE_OF_RECOMMENDATION"]) for row in diagnostic_rows)
    summary = [
        {"METRIC": "TOTAL_DIVERGENCES", "VALUE": len(diagnostic_rows)},
        {"METRIC": "FALSE_POSITIVES", "VALUE": len(false_positives)},
        {"METRIC": "FALSE_NEGATIVES", "VALUE": len(false_negatives)},
        {"METRIC": "UNCLEAR_CURRENT", "VALUE": sum(row.get("MESSAGE_OUTCOME") == "UNCLEAR" for row in replay_rows)},
        {"METRIC": "HISTORICAL_ONLY_DIVERGENCES", "VALUE": len(temporal)},
        {"METRIC": "CASE_STATE_DIVERGENCES", "VALUE": len(case_state)},
        {"METRIC": "ORGANIZATION_ONLY_DIFFERENCES", "VALUE": len(organization)},
        {"METRIC": "HUMAN_SAMPLE_SIZE", "VALUE": min(150, len(diagnostic_rows))},
    ]
    root_summary = [{"ROOT_CAUSE": key, "COUNT": value} for key, value in root_counts.most_common()]
    rule_recommendations = [{"CURRENT_RULE_ID": key[0], "RECOMMENDATION": key[1], "CONFIDENCE_OF_RECOMMENDATION": key[2], "COUNT": count} for key, count in recommendations.most_common()]
    fp_groups = _group_rows(false_positives, ("CURRENT_RULE_ID", "CURRENT_FAMILY", "CURRENT_EVENT_TYPE", "CURRENT_ACTION_TYPE", "CURRENT_SCOPE", "REFERENCE_MESSAGE_OUTCOME", "REFERENCE_ACTION_REQUIRED", "REFERENCE_ACTION_TYPE", "REFERENCE_FAMILY", "REFERENCE_EVENT", "REFERENCE_SCOPE"))
    fn_groups = _group_rows(false_negatives, ("REFERENCE_RULE_ID", "REFERENCE_FAMILY", "REFERENCE_EVENT", "REFERENCE_ACTION_TYPE", "CURRENT_MESSAGE_OUTCOME", "CURRENT_RULE_ID"))
    unclear_groups = _group_rows(unclear, ("subject", "from_email", "CURRENT_FAMILY", "REFERENCE_MESSAGE_OUTCOME", "REFERENCE_ACTION_TYPE", "REFERENCE_SCOPE", "UNCLEAR_REFERENCE_BUCKET"))
    for row in signature:
        row["SIGNATURE_SUBGROUP"] = _signature_subgroup(row)
    for row in payment:
        row["PAYMENT_SUBGROUP"] = _payment_subgroup(row)
    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_sheet(workbook, "ExecutiveSummary", summary)
    _write_sheet(workbook, "AllDivergences", diagnostic_rows)
    _write_sheet(workbook, "FalsePositives", fp_groups + false_positives)
    _write_sheet(workbook, "FalseNegatives", fn_groups + false_negatives)
    _write_sheet(workbook, "SignatureAudit", signature)
    _write_sheet(workbook, "PaymentAudit", payment)
    _write_sheet(workbook, "UnclearAudit", unclear_groups + unclear)
    _write_sheet(workbook, "TemporalityAudit", temporal)
    _write_sheet(workbook, "CaseStateAudit", case_state)
    _write_sheet(workbook, "OrganizationAudit", organization)
    _write_sheet(workbook, "RootCauseSummary", root_summary)
    _write_sheet(workbook, "RuleRecommendations", rule_recommendations)
    _write_sheet(workbook, "HumanValidationSample", _sample([dict(row) for row in diagnostic_rows]))
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return {
        "total_divergences": len(diagnostic_rows),
        "false_positives": false_positives,
        "false_negatives": false_negatives,
        "unclear": unclear,
        "root_counts": root_counts,
        "recommendations": rule_recommendations,
        "signature": signature,
        "payment": payment,
        "temporal": temporal,
        "case_state": case_state,
        "organization": organization,
        "output_path": str(output_path),
    }


def _signature_subgroup(row):
    subject = str(row.get("subject") or "").lower()
    body = str(row.get("body_preview") or "").lower()
    text = f"{subject} {body}"
    if "rechaz" in text:
        return "SIGNATURE_REJECTED"
    if "expir" in text:
        return "SIGNATURE_EXPIRED"
    if "completad" in text or "exitos" in text:
        return "SIGNATURE_SUCCESS"
    if "error" in text:
        return "SIGNATURE_ERROR"
    if "notific" in text or "inform" in text:
        return "INFORMATIONAL_NOTIFICATION"
    return "AMBIGUOUS"


def _payment_subgroup(row):
    text = f"{row.get('subject', '')} {row.get('body_preview', '')}".lower()
    if "no legal" in text or "legaliz" in text or "legalizar" in text:
        return "LEGALIZATION_REQUEST_OR_FAILURE"
    if "aplicado correctamente" in text or "pago aplicado" in text:
        return "PAYMENT_APPLIED"
    if "comprobante" in text:
        return "NORMAL_RECEIPT"
    if "recibo" in text:
        return "RECEIPT"
    if "inform" in text:
        return "INFORMATIONAL"
    return "AMBIGUOUS"
