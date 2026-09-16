import csv
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace

from django.test import TestCase

from email_exceptions.correlation import FunctionalReferences, correlate_email_to_case, extract_functional_references, _case_references
from email_exceptions.correlation_replay import ReplayCase, _correlate_with_precomputed_features, _false_split_rows, run_correlation_replay
from email_exceptions.models import CaseMessage, EmailException, ExceptionCase, InboundEmail


TIMING_METRICS = {
    "CLASSIFICATION_SECONDS",
    "CORRELATION_SECONDS",
    "DIAGNOSTICS_SECONDS",
    "EXCEL_WRITE_SECONDS",
    "TOTAL_SECONDS",
}


class CorrelationReplayTests(TestCase):
    fields = ("entry_id", "message_id", "message_id_hash", "conversation_id", "received_at", "from_name", "from_email", "from_domain", "subject_original", "subject_normalized", "body_text", "body_preview", "has_attachments", "attachment_count")

    def write_fixture(self, directory):
        path = Path(directory) / "emails.csv"
        rows = [
            {"entry_id": "5", "message_id": "m5", "message_id_hash": "h5", "received_at": "2026-01-05T10:00:00+00:00", "subject_original": "Consulta general", "body_text": "Necesito confirmar el estado."},
            {"entry_id": "2", "message_id": "m2", "message_id_hash": "h2", "received_at": "2026-01-02T10:00:00+00:00", "subject_original": "Boletín informativo", "body_text": "Información general."},
            {"entry_id": "1", "message_id": "m1", "message_id_hash": "h1", "received_at": "2026-01-01T10:00:00+00:00", "subject_original": "Pago duplicado reclamación ABC-123", "body_text": "Requiere conciliación."},
            {"entry_id": "4", "message_id": "m4", "message_id_hash": "h4", "received_at": "2026-01-04T10:00:00+00:00", "subject_original": "Listado de 120 pólizas", "body_text": "Inconsistencia masiva."},
            {"entry_id": "3", "message_id": "m3", "message_id_hash": "h3", "received_at": "2026-01-03T10:00:00+00:00", "subject_original": "Pago aplicado correctamente reclamación ABC-123", "body_text": "Operación procesada correctamente."},
        ]
        with path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=self.fields)
            writer.writeheader()
            for row in rows:
                writer.writerow(row)
        return path

    def test_replay_orders_messages_chronologically_and_is_deterministic(self):
        with TemporaryDirectory() as directory:
            dataset = self.write_fixture(directory)
            first = run_correlation_replay(dataset, Path(directory) / "first.xlsx")
            second = run_correlation_replay(dataset, Path(directory) / "second.xlsx")
        self.assertEqual([row["message_id_hash"] for row in first["decisions"]], ["h1", "h2", "h3", "h4", "h5"])
        self.assertEqual(first["decisions"], second["decisions"])

        def case_snapshot(case):
            return {
                "case_key": case.case_key,
                "organization": case.organization,
                "family": case.family,
                "event_type": case.event_type,
                "status": case.status,
                "action_type": case.action_type,
                "scope": case.scope,
                "opened_at": case.opened_at,
                "last_activity_at": case.last_activity_at,
                "resolved_at": case.resolved_at,
                "functional_references": case.functional_references,
                "conversation_ids": tuple(sorted(case.conversation_ids)),
                "messages": tuple(getattr(link.email, "external_message_id", "") for link in case.messages),
            }

        self.assertEqual([case_snapshot(case) for case in first["cases"]], [case_snapshot(case) for case in second["cases"]])
        self.assertEqual(first["potential_resolutions"], second["potential_resolutions"])
        self.assertEqual(first["potential_false_merges"], second["potential_false_merges"])
        self.assertEqual(first["potential_false_splits"], second["potential_false_splits"])

        first_metrics = dict(first["metrics"])
        second_metrics = dict(second["metrics"])
        for metric in TIMING_METRICS:
            self.assertIn(metric, first_metrics)
            self.assertIn(metric, second_metrics)
            self.assertIsInstance(first_metrics[metric], (int, float))
            self.assertIsInstance(second_metrics[metric], (int, float))
            self.assertGreaterEqual(first_metrics[metric], 0)
            self.assertGreaterEqual(second_metrics[metric], 0)
            first_metrics.pop(metric)
            second_metrics.pop(metric)
        self.assertEqual(first_metrics, second_metrics)

    def test_origin_policy_does_not_create_cases_for_noise_info_unclear_or_scope(self):
        with TemporaryDirectory() as directory:
            result = run_correlation_replay(self.write_fixture(directory), Path(directory) / "replay.xlsx")
        metrics = result["metrics"]
        self.assertEqual(metrics["SIMULATED_CASES"], 1)
        self.assertEqual(metrics["BATCH_MESSAGES"], 1)
        self.assertEqual(metrics["SUCCESS_LINKED_TO_CASE"], 1)
        self.assertEqual(metrics.get("UNCLEAR_LINKED_TO_CASE", 0), 0)
        self.assertEqual(metrics.get("INFORMATIONAL_LINKED_TO_CASE", 0), 0)

    def test_replay_is_read_only_for_application_models(self):
        with TemporaryDirectory() as directory:
            before = (ExceptionCase.objects.count(), CaseMessage.objects.count(), EmailException.objects.count(), InboundEmail.objects.count())
            run_correlation_replay(self.write_fixture(directory), Path(directory) / "replay.xlsx")
            after = (ExceptionCase.objects.count(), CaseMessage.objects.count(), EmailException.objects.count(), InboundEmail.objects.count())
        self.assertEqual(before, after)

    def test_success_is_marked_as_potential_resolution_without_closing_case(self):
        with TemporaryDirectory() as directory:
            result = run_correlation_replay(self.write_fixture(directory), Path(directory) / "replay.xlsx")
        resolutions = result["potential_resolutions"]
        self.assertEqual(len(resolutions), 1)
        self.assertTrue(resolutions[0]["POTENTIAL_RESOLUTION"])
        self.assertEqual(result["cases"][0].status, "OPEN")

    def test_false_split_diagnostic_does_not_modify_cases(self):
        message_a = SimpleNamespace(subject="Reclamación ABC-123", body_text="Seguimiento del caso.", conversation_id="conversation-a")
        message_b = SimpleNamespace(subject="Reclamación ABC-123", body_text="Información adicional.", conversation_id="conversation-b")
        cases = [
            ReplayCase("claim:ABC-123", functional_references=FunctionalReferences(claim_number="ABC-123"), messages=[SimpleNamespace(email=message_a)], conversation_ids={"conversation-a"}),
            ReplayCase("claim:ABC-123|policy:M-1", functional_references=FunctionalReferences(claim_number="ABC-123"), messages=[SimpleNamespace(email=message_b)], conversation_ids={"conversation-b"}),
        ]
        before = [(case.case_key, list(case.messages), case.functional_references, set(case.conversation_ids), case.status, case.organization, case.family, case.event_type, case.action_type, case.scope) for case in cases]
        rows = _false_split_rows(cases)
        self.assertTrue(rows)
        after = [(case.case_key, list(case.messages), case.functional_references, set(case.conversation_ids), case.status, case.organization, case.family, case.event_type, case.action_type, case.scope) for case in cases]
        self.assertEqual(after, before)

    def test_report_has_requested_sheets_and_empty_human_columns(self):
        with TemporaryDirectory() as directory:
            output = Path(directory) / "replay.xlsx"
            result = run_correlation_replay(self.write_fixture(directory), output)
            from openpyxl import load_workbook
            workbook = load_workbook(output, read_only=True, data_only=True)
            self.assertIn("HumanValidationSample", workbook.sheetnames)
            sheet = workbook["HumanValidationSample"]
            headers = [cell.value for cell in next(sheet.iter_rows())]
            rows = list(sheet.iter_rows(min_row=2, values_only=True))
            workbook.close()
            for column in ("VALIDACION_HUMANA", "COMENTARIO_HUMANO"):
                if column in headers:
                    index = headers.index(column)
                    self.assertTrue(all(row[index] in (None, "") for row in rows))
        self.assertTrue(result["output_path"].endswith("replay.xlsx"))

    def test_replay_precomputes_email_features_once_per_message(self):
        with TemporaryDirectory() as directory:
            result = run_correlation_replay(self.write_fixture(directory), Path(directory) / "replay.xlsx")
        self.assertEqual(result["metrics"]["FEATURE_EXTRACTION_CALLS"], result["metrics"]["TOTAL_MESSAGES"])
        self.assertGreater(result["metrics"]["FEATURE_EXTRACTION_CALLS_ESTIMATED_BEFORE"], result["metrics"]["FEATURE_EXTRACTION_CALLS"])
        self.assertLessEqual(result["metrics"]["CASE_FEATURE_EXTRACTION_CALLS"], result["metrics"]["CASE_FEATURE_EXTRACTION_CALLS_ESTIMATED_BEFORE"])

    def test_precomputed_path_matches_public_engine_decisions(self):
        scenarios = [
            ("Reclamación ABC-123", "", "", "FAILURE", "CASE", [ReplayCase("claim:ABC-123")]),
            ("Seguimiento", "", "conv-1", "FAILURE", "CASE", [ReplayCase("case:1", messages=[SimpleNamespace(email=SimpleNamespace(conversation_id="conv-1", subject="Anterior", body_text=""))])]),
            ("Reclamación ABC-123", "", "", "FAILURE", "CASE", [ReplayCase("claim:XYZ-999")]),
            ("Seguimiento", "", "conv-1", "FAILURE", "CASE", [ReplayCase("case:1", messages=[SimpleNamespace(email=SimpleNamespace(conversation_id="conv-1", subject="Anterior", body_text=""))]), ReplayCase("case:2", messages=[SimpleNamespace(email=SimpleNamespace(conversation_id="conv-1", subject="Anterior", body_text=""))])]),
            ("Póliza M-600000598 periodo 2026-08", "", "", "FAILURE", "CASE", [ReplayCase("policy:M-600000598|period:2026-08", organization="BEMSA", family="PAGO", event_type="PAGO_DOBLE")]),
            ("Consulta general", "", "", "FAILURE", "CASE", [ReplayCase("claim:ABC-123")]),
            ("Reclamación ABC-123", "", "", "SUCCESS", "CASE", [ReplayCase("claim:ABC-123")]),
            ("Reclamación ABC-123", "", "", "FAILURE", "BATCH", [ReplayCase("claim:ABC-123")]),
            ("Reclamación ABC-123", "", "", "FAILURE", "PLATFORM", [ReplayCase("claim:ABC-123")]),
        ]
        for subject, body, conversation_id, outcome, scope, candidates in scenarios:
            with self.subTest(subject=subject, outcome=outcome, scope=scope):
                email = InboundEmail(subject=subject, body_text=body, conversation_id=conversation_id)
                classification = SimpleNamespace(message_outcome=outcome, scope=scope, organization="", family="", event_type="")
                references = extract_functional_references(email)
                case_references = {id(case): _case_references(case) for case in candidates}
                expected = correlate_email_to_case(email, classification, candidates)
                actual = _correlate_with_precomputed_features(email, classification, candidates, references, case_references)
                self.assertEqual((actual.matched, actual.case, actual.method, actual.confidence, actual.reason, actual.evidence, actual.conflicts), (expected.matched, expected.case, expected.method, expected.confidence, expected.reason, expected.evidence, expected.conflicts))
