import csv
import tempfile
from pathlib import Path

from django.core import management
from django.core.management.base import CommandError
from django.test import TestCase, override_settings

from email_exceptions.backfill import row_to_payload, run_backfill
from email_exceptions.management.commands.email_exceptions_backfill import validate_isolated_database
from email_exceptions.models import CaseMessage, EmailException, ExceptionCase, InboundEmail


class BackfillTests(TestCase):
    def row(self, message_id, received_at, subject="Comprobante de pago BEMSA"):
        return {
            "entry_id": message_id,
            "message_id": message_id,
            "message_id_hash": "hash-" + message_id,
            "conversation_id": "conv-1",
            "received_at": received_at,
            "from_name": "Pagos",
            "from_email": "pagos@bemsa.com.co",
            "subject_original": subject,
            "body_text": "Adjunto comprobante de pago de la póliza BEMSA.",
            "source_store": "comunicaciones@segurosays.com",
            "source_folder": "Bandeja de entrada",
        }

    def csv_file(self, rows):
        handle = tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", newline="", suffix=".csv", delete=False)
        writer = csv.DictWriter(handle, fieldnames=sorted(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
        handle.close()
        self.addCleanup(lambda: Path(handle.name).unlink(missing_ok=True))
        return Path(handle.name)

    def test_mapping_uses_source_store_and_stable_message_id(self):
        payload = row_to_payload(self.row("message-1", "2026-01-01T01:00:00+00:00"))
        self.assertEqual(payload["source_mailbox"], "comunicaciones@segurosays.com")
        self.assertEqual(payload["message_id"], "message-1")
        self.assertEqual(payload["conversation_id"], "conv-1")

    def test_rows_are_processed_chronologically_with_stable_tiebreak(self):
        path = self.csv_file([
            self.row("z", "2026-01-01T02:00:00+00:00"),
            self.row("a", "2026-01-01T01:00:00+00:00"),
        ])
        ordered = __import__("email_exceptions.backfill", fromlist=["_ordered_rows"])._ordered_rows(path)
        self.assertEqual([row["message_id"] for row in ordered], ["a", "z"])

    def test_dry_run_rolls_back_all_objects(self):
        path = self.csv_file([self.row("dry-1", "2026-01-01T01:00:00+00:00")])
        with tempfile.TemporaryDirectory() as directory:
            result = run_backfill(path, Path(directory) / "report.xlsx", limit=100)
        self.assertEqual(result["before"], result["after"])
        self.assertEqual(InboundEmail.objects.count(), 0)
        self.assertEqual(ExceptionCase.objects.count(), 0)
        self.assertEqual(CaseMessage.objects.count(), 0)
        self.assertEqual(EmailException.objects.count(), 0)

    def test_command_rejects_persistent_default_database(self):
        path = self.csv_file([self.row("persist-1", "2026-01-01T01:00:00+00:00")])
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(CommandError):
                management.call_command("email_exceptions_backfill", dataset=path, output=Path(directory) / "report.xlsx", persist=True)

    def test_command_rejects_unknown_database_alias(self):
        path = self.csv_file([self.row("unknown-1", "2026-01-01T01:00:00+00:00")])
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(CommandError):
                management.call_command("email_exceptions_backfill", dataset=path, output=Path(directory) / "report.xlsx", database="missing")

    @override_settings(DATABASES={
        "default": {"ENGINE": "django.db.backends.postgresql", "NAME": "ays_vault", "HOST": "db", "PORT": "5432", "USER": "vault"},
        "backfill": {"ENGINE": "django.db.backends.postgresql", "NAME": "ays_vault", "HOST": "db", "PORT": "5432", "USER": "backfill"},
    })
    def test_isolation_rejects_same_physical_database_even_with_different_user(self):
        with self.assertRaises(CommandError) as raised:
            validate_isolated_database("backfill")
        self.assertEqual(str(raised.exception), "BACKFILL_DATABASE_MUST_BE_ISOLATED")

    def test_report_contains_required_sheets(self):
        path = self.csv_file([self.row("report-1", "2026-01-01T01:00:00+00:00")])
        with tempfile.TemporaryDirectory() as directory:
            result = run_backfill(path, Path(directory) / "report.xlsx", limit=1)
            from openpyxl import load_workbook
            workbook = load_workbook(result["output"], read_only=True)
            self.assertEqual(workbook.sheetnames, ["Summary", "Outcomes", "Cases", "HighMatches", "ConflictsAmbiguities", "Errors", "InvariantAudit", "Idempotency"])
            workbook.close()

    def test_missing_source_store_is_reported_as_data_error(self):
        row = self.row("invalid-1", "2026-01-01T01:00:00+00:00")
        row["source_store"] = "Archive"
        path = self.csv_file([row])
        with tempfile.TemporaryDirectory() as directory:
            result = run_backfill(path, Path(directory) / "report.xlsx", limit=1)
        self.assertEqual(result["summary"]["DATA_ERROR"], 1)
        self.assertEqual(result["summary"]["MESSAGES_PROCESSED"], 0)
