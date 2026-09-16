from __future__ import annotations

import csv
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .models import InboundEmail
from .services import classify_inbound_email


csv.field_size_limit(2**31 - 1)


STRUCTURED_FIELDS = {
    "message_id", "message_id_hash", "conversation_id", "entry_id",
    "case_key", "correlation_key", "CURRENT_MESSAGE_IS_EXCEPTION",
    "ELIGIBLE_FOR_EMAIL_EXCEPTION",
}


def _bool(value):
    return str(value or "").strip().lower() in {"1", "true", "yes", "si", "sí"}


def _text(value, limit=32000):
    value = "" if value is None else str(value)
    return value if len(value) <= limit else value[:limit] + "…"


def _historical_email(row):
    return InboundEmail(
        source_mailbox="aysltda@asesorsura.com",
        external_message_id=row.get("message_id") or row.get("entry_id") or row.get("message_id_hash") or "historical",
        conversation_id=row.get("conversation_id") or "",
        received_at=row.get("received_at") or "",
        from_name=row.get("from_name") or "",
        from_email=row.get("from_email") or "",
        from_domain=row.get("from_domain") or "",
        subject=row.get("subject_original") or row.get("subject_normalized") or "",
        body_text=row.get("body_text") or row.get("body_preview") or "",
        has_attachments=_bool(row.get("has_attachments")) or int(row.get("attachment_count") or 0) > 0,
    )


def _reference_rows(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Cambios"]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "") for value in next(rows)]
        return {
            str(row[headers.index("message_id_hash")] or ""): dict(zip(headers, row))
            for row in rows
            if len(row) > headers.index("message_id_hash") and row[headers.index("message_id_hash")]
        }
    finally:
        workbook.close()


def _reference_case_count(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return max(0, workbook["QueueCases"].max_row - 1)
    finally:
        workbook.close()


def _current_row(row, result, reference):
    current_exception = result.current_message_is_exception
    eligible = result.eligible_for_email_exception
    reference_queue = _bool(reference.get("V631_SHOULD_BE_IN_EXCEPTION_QUEUE"))
    current_queue = eligible
    if reference_queue and current_queue:
        queue_result = "TRUE_POSITIVE"
    elif not reference_queue and not current_queue:
        queue_result = "TRUE_NEGATIVE"
    elif current_queue:
        queue_result = "POTENTIAL_FALSE_POSITIVE"
    else:
        queue_result = "POTENTIAL_FALSE_NEGATIVE"
    comparisons = {
        "OUTCOME_MATCH": str(reference.get("V631_MESSAGE_OUTCOME") or "") == result.message_outcome,
        "ACTION_REQUIRED_MATCH": _bool(reference.get("V631_ACTION_REQUIRED")) == result.action_required,
        "ACTION_TYPE_MATCH": str(reference.get("V631_ACTION_TYPE") or "") == result.action_type,
        "SCOPE_MATCH": str(reference.get("V631_SCOPE") or "") == result.scope,
        "QUEUE_MATCH": reference_queue == current_queue,
    }
    optional_comparisons = {
        "FAMILY_MATCH": ("V631_FAMILY", result.family),
        "EVENT_MATCH": ("V631_EVENT", result.event_type),
        "ORGANIZATION_MATCH": ("V631_ORGANIZATION", result.organization),
    }
    for name, (reference_field, current_value) in optional_comparisons.items():
        if reference.get(reference_field) not in (None, ""):
            comparisons[name] = str(reference[reference_field]) == current_value
    return {
        "message_id": row.get("message_id", ""),
        "message_id_hash": row.get("message_id_hash", ""),
        "entry_id": row.get("entry_id", ""),
        "conversation_id": row.get("conversation_id", ""),
        "received_at": row.get("received_at", ""),
        "sender": row.get("from_email", ""),
        "subject": row.get("subject_original") or row.get("subject_normalized", ""),
        "body_text": _text(row.get("body_text") or row.get("body_preview", "")),
        "MESSAGE_OUTCOME": result.message_outcome,
        "ACTION_REQUIRED": result.action_required,
        "ACTION_TYPE": result.action_type,
        "SCOPE": result.scope,
        "ORGANIZATION": result.organization,
        "FAMILY": result.family,
        "EVENT_TYPE": result.event_type,
        "RULE_ID": result.rule_id,
        "CONFIDENCE": result.confidence,
        "REASON": result.exception_reason,
        "CORRELATION_KEY": result.correlation_key,
        "CURRENT_MESSAGE_IS_EXCEPTION": current_exception,
        "ELIGIBLE_FOR_EMAIL_EXCEPTION": eligible,
        "V631_MESSAGE_OUTCOME": reference.get("V631_MESSAGE_OUTCOME", ""),
        "V631_ACTION_REQUIRED": reference.get("V631_ACTION_REQUIRED", ""),
        "V631_ACTION_TYPE": reference.get("V631_ACTION_TYPE", ""),
        "V631_SCOPE": reference.get("V631_SCOPE", ""),
        "V631_FAMILY": reference.get("V631_FAMILY", ""),
        "V631_EVENT": reference.get("V631_EVENT", ""),
        "V631_ORGANIZATION": reference.get("V631_ORGANIZATION", ""),
        "V631_QUEUE": reference_queue,
        "CURRENT_QUEUE": current_queue,
        "QUEUE_RESULT": queue_result,
        **comparisons,
        "OVERALL_MATCH": all(comparisons.values()),
    }


def _write_sheet(workbook, title, rows):
    sheet = workbook.create_sheet(title)
    if not rows:
        sheet.append(["No data"])
        return
    headers = list(rows[0])
    sheet.append(headers)
    for row in rows:
        sheet.append([_text(row.get(header)) if header not in STRUCTURED_FIELDS else row.get(header) for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def run_replay(dataset_path, reference_path, output_path):
    reference = _reference_rows(reference_path)
    current_rows, comparisons = [], []
    with Path(dataset_path).open("r", encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            result = classify_inbound_email(_historical_email(row))
            reference_row = reference.get(row.get("message_id_hash", ""), {})
            current = _current_row(row, result, reference_row)
            current["REFERENCE_VALIDATION_SAMPLE"] = "NOT_LINKED_NO_STABLE_KEY"
            current_rows.append(current)
            if reference_row:
                comparisons.append(current)

    outcome_counts = Counter(row["MESSAGE_OUTCOME"] for row in current_rows)
    action_counts = Counter(row["ACTION_TYPE"] for row in current_rows)
    scope_counts = Counter(row["SCOPE"] for row in current_rows)
    rule_counts = Counter(row["RULE_ID"] for row in current_rows)
    queue_counts = Counter(row["QUEUE_RESULT"] for row in comparisons)
    batch_violations = [row for row in current_rows if row["SCOPE"] == "BATCH" and row["ELIGIBLE_FOR_EMAIL_EXCEPTION"]]
    platform_violations = [row for row in current_rows if row["SCOPE"] == "PLATFORM" and row["ELIGIBLE_FOR_EMAIL_EXCEPTION"]]
    unclear_rows = [row for row in comparisons if row["V631_MESSAGE_OUTCOME"] == "UNCLEAR" or row["MESSAGE_OUTCOME"] == "UNCLEAR"]
    protected_terms = ("sharefile", "boletin", "newsletter", "reunion", "recibo", "factura", "pago aplicado", "firma completada", "renovado")
    protected_rows = [row for row in current_rows if any(term in row["subject"].lower() for term in protected_terms)]

    summary = [
        {"METRIC": "TOTAL_MESSAGES", "VALUE": len(current_rows)},
        {"METRIC": "TOTAL_REFERENCE_CASES", "VALUE": _reference_case_count(reference_path)},
        {"METRIC": "CURRENT_ELIGIBLE_MESSAGES", "VALUE": sum(row["ELIGIBLE_FOR_EMAIL_EXCEPTION"] for row in current_rows)},
        {"METRIC": "REFERENCE_COMPARABLE_MESSAGES", "VALUE": len(comparisons)},
        {"METRIC": "QUEUE_TRUE_POSITIVE", "VALUE": queue_counts["TRUE_POSITIVE"]},
        {"METRIC": "QUEUE_TRUE_NEGATIVE", "VALUE": queue_counts["TRUE_NEGATIVE"]},
        {"METRIC": "POTENTIAL_FALSE_POSITIVE", "VALUE": queue_counts["POTENTIAL_FALSE_POSITIVE"]},
        {"METRIC": "POTENTIAL_FALSE_NEGATIVE", "VALUE": queue_counts["POTENTIAL_FALSE_NEGATIVE"]},
        {"METRIC": "QUEUE_AGREEMENT_PERCENT", "VALUE": round(100 * (queue_counts["TRUE_POSITIVE"] + queue_counts["TRUE_NEGATIVE"]) / len(comparisons), 2) if comparisons else None},
        {"METRIC": "OUTCOME_AGREEMENT_PERCENT", "VALUE": round(100 * sum(row["OUTCOME_MATCH"] for row in comparisons) / len(comparisons), 2) if comparisons else None},
        {"METRIC": "ACTION_REQUIRED_AGREEMENT_PERCENT", "VALUE": round(100 * sum(row["ACTION_REQUIRED_MATCH"] for row in comparisons) / len(comparisons), 2) if comparisons else None},
        {"METRIC": "ACTION_TYPE_AGREEMENT_PERCENT", "VALUE": round(100 * sum(row["ACTION_TYPE_MATCH"] for row in comparisons) / len(comparisons), 2) if comparisons else None},
        {"METRIC": "SCOPE_AGREEMENT_PERCENT", "VALUE": round(100 * sum(row["SCOPE_MATCH"] for row in comparisons) / len(comparisons), 2) if comparisons else None},
        {"METRIC": "BATCH_ELIGIBLE_VIOLATIONS", "VALUE": len(batch_violations)},
        {"METRIC": "PLATFORM_ELIGIBLE_VIOLATIONS", "VALUE": len(platform_violations)},
    ]
    distributions = []
    for name, counter in (("OUTCOME", outcome_counts), ("ACTION_TYPE", action_counts), ("SCOPE", scope_counts), ("RULE_ID", rule_counts)):
        distributions.extend({"DIMENSION": name, "VALUE": key, "COUNT": value} for key, value in counter.most_common())

    by_stratum = {}
    for row in current_rows:
        if row in comparisons:
            stratum = row["QUEUE_RESULT"]
        elif row["MESSAGE_OUTCOME"] == "UNCLEAR":
            stratum = "UNCLEAR"
        else:
            stratum = row["MESSAGE_OUTCOME"]
        by_stratum.setdefault(stratum, []).append(row)
    human_sample = []
    for stratum in sorted(by_stratum):
        for row in by_stratum[stratum][:5]:
            human_sample.append({"STRATUM": stratum, **row, "VALIDACION_HUMANA": "", "COMENTARIO_HUMANO": ""})

    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_sheet(workbook, "Resumen", summary)
    _write_sheet(workbook, "ClasificacionActual", current_rows)
    _write_sheet(workbook, "ComparacionV631", comparisons)
    _write_sheet(workbook, "QueueDivergences", [row for row in comparisons if not row["QUEUE_MATCH"]])
    _write_sheet(workbook, "PotentialFalsePositives", [row for row in comparisons if row["QUEUE_RESULT"] == "POTENTIAL_FALSE_POSITIVE"])
    _write_sheet(workbook, "PotentialFalseNegatives", [row for row in comparisons if row["QUEUE_RESULT"] == "POTENTIAL_FALSE_NEGATIVE"])
    _write_sheet(workbook, "BatchPlatformAudit", batch_violations + platform_violations)
    _write_sheet(workbook, "UnclearAudit", unclear_rows)
    _write_sheet(workbook, "ProtectedNegatives", protected_rows)
    _write_sheet(workbook, "RulesDistribution", distributions)
    _write_sheet(workbook, "HumanValidationSample", human_sample)
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return {
        "total_messages": len(current_rows),
        "current_rows": current_rows,
        "comparisons": comparisons,
        "summary": summary,
        "queue_counts": queue_counts,
        "batch_violations": batch_violations,
        "platform_violations": platform_violations,
        "output_path": str(output_path),
    }
