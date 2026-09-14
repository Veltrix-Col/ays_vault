"""Puente entre los archivos subidos (Django) y el motor `gestor_cotizaciones`.

Mismo patrón que `conciliacion.services.processor`: se recibe el contenido ya
validado por el formulario, se materializa en un directorio temporal, se corre
el motor (registro -> adaptadores -> análisis -> render) y se devuelve el xlsx
+ un resumen serializable. El Excel se endurece (anti fórmula, sin
hipervínculos) antes de salir, igual que en `conciliacion`.

A diferencia del CLI (`gestor_cotizaciones.cli`), que deja JSON canónicos
intermedios en `out/canonico/` para que el analista los edite a mano, aquí todo
corre en memoria en una sola pasada: la app web no tiene ese flujo de edición.
"""
from __future__ import annotations

import importlib
import io
import logging
import os
import re
from dataclasses import dataclass
from pathlib import Path
from tempfile import TemporaryDirectory
from time import monotonic

from django.conf import settings
from openpyxl import load_workbook

from gestor_cotizaciones.core.adaptador import cargar_catalogo
from gestor_cotizaciones.core.analisis import analizar
from gestor_cotizaciones.core.llm import LLM
from gestor_cotizaciones.core.registry import asignar_archivos, cargar_registro, instanciar

logger = logging.getLogger("cotizador")

PKG_DIR = Path(settings.BASE_DIR) / "gestor_cotizaciones"
CONFIG_DIR = PKG_DIR / "config"
CACHE_DIR = Path(settings.BASE_DIR) / "runtime" / "gestor_cotizaciones" / "cache_llm"
PROMPT_DIR = Path(settings.BASE_DIR) / "runtime" / "gestor_cotizaciones" / "prompts"


class CotizadorProcessingError(ValueError):
    """Error de negocio al generar el consolidado; una vista puede mostrar el mensaje tal cual."""


_NOMBRE_SEGURO = re.compile(r"[^A-Za-z0-9 ._()\-À-ſ]")


def _nombre_seguro(nombre: str, respaldo: str) -> str:
    base = os.path.basename((nombre or "").replace("\\", "/"))
    base = _NOMBRE_SEGURO.sub("_", base).strip().strip(".")
    return base or respaldo


def _volcar(archivo, destino: Path) -> Path:
    archivo.seek(0)
    with destino.open("wb") as stream:
        for chunk in archivo.chunks():
            stream.write(chunk)
    archivo.seek(0)
    return destino


def _neutralizar(valor):
    if isinstance(valor, str) and valor[:1] in {"=", "+", "-", "@"}:
        return "'" + valor
    return valor


def _endurecer_xlsx(contenido: bytes) -> bytes:
    """Neutraliza celdas de texto con prefijo de fórmula y quita hipervínculos
    (mismo criterio que `conciliacion.services.processor._endurecer_xlsx`)."""
    workbook = load_workbook(io.BytesIO(contenido))
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    cell.hyperlink = None
                    if isinstance(cell.value, str):
                        cell.value = _neutralizar(cell.value)
        salida = io.BytesIO()
        workbook.save(salida)
        return salida.getvalue()
    finally:
        workbook.close()


@dataclass(frozen=True)
class CotizadorOutput:
    content: bytes
    filename: str
    summary: dict[str, object]


def generar_consolidado(*, ramo: str, nit: str, actual: str, otra_aseguradora: str,
                         archivos: list) -> CotizadorOutput:
    """`archivos` son los `UploadedFile` ya validados por `CotizadorUploadForm`
    (cualquier mezcla de aseguradoras/formatos). `actual` es el id de la
    aseguradora vigente (aporta la prima anterior al análisis de flota)."""
    started = monotonic()
    catalogo = cargar_catalogo(ramo, base=CONFIG_DIR)
    registro = cargar_registro(ramo, base=CONFIG_DIR)
    llm = LLM(cache_dir=CACHE_DIR, prompt_dir=PROMPT_DIR)

    with TemporaryDirectory(prefix="ays-cotizador-") as directorio:
        raiz = Path(directorio)
        rutas = [
            _volcar(archivo, raiz / _nombre_seguro(getattr(archivo, "name", ""), f"archivo_{i}"))
            for i, archivo in enumerate(archivos)
        ]

        grupos, sueltos = asignar_archivos(registro, rutas)
        if sueltos and not otra_aseguradora:
            nombres = ", ".join(p.name for p in sueltos)
            raise CotizadorProcessingError(
                f"No reconozco estos archivos como ninguna aseguradora configurada: {nombres}. "
                "Indique un nombre en «Nombre para archivos no reconocidos» para procesarlos con "
                "el adaptador genérico, o quítelos de la carga."
            )
        if sueltos:
            grupos[otra_aseguradora] = sueltos

        ids_conocidos = {a["id"] for a in registro["adaptadores"]}
        advertencias: list[str] = []
        cotizaciones = []
        for aseguradora_id, files in grupos.items():
            es_actual = aseguradora_id == actual
            try:
                if aseguradora_id in ids_conocidos:
                    adaptador = instanciar(registro, aseguradora_id, catalogo=catalogo, llm=llm,
                                            es_actual=es_actual, opciones={"nit": nit})
                else:
                    adaptador = instanciar(
                        registro, "generico", catalogo=catalogo, llm=llm, es_actual=es_actual,
                        opciones={"nit": nit, "id": aseguradora_id,
                                  "nombre": aseguradora_id.upper(), "ramo": ramo},
                    )
                    advertencias.append(
                        f"{aseguradora_id}: procesado con el adaptador genérico (100% IA, sin parser propio)."
                    )
                cotizaciones.append(adaptador.extraer(files))
            except CotizadorProcessingError:
                raise
            except Exception as exc:
                logger.exception("Fallo al extraer los archivos de %s", aseguradora_id)
                raise CotizadorProcessingError(
                    f"No fue posible leer los archivos de «{aseguradora_id}»: {exc}"
                ) from exc

        if not any(c.es_actual for c in cotizaciones):
            raise CotizadorProcessingError(
                f"Ninguno de los archivos cargados corresponde a la compañía actual indicada ({actual})."
            )

        cotizaciones.sort(key=lambda c: (not c.es_actual, c.aseguradora))
        actual_cot = next(c for c in cotizaciones if c.es_actual)
        placas = [v.placa for v in actual_cot.vehiculos]
        analisis = analizar(cotizaciones, catalogo, set(placas), llm)

        render_mod = importlib.import_module(f"gestor_cotizaciones.ramos.{ramo}.render")
        render_cls = next(v for k, v in vars(render_mod).items() if k.startswith("Render"))

        with TemporaryDirectory(prefix="ays-cotizador-out-") as out_dir:
            salida = Path(out_dir) / "consolidado.xlsx"
            render_cls(catalogo, cotizaciones, analisis, placas).generar(salida)
            contenido = _endurecer_xlsx(salida.read_bytes())

        summary = {
            "ramo": ramo,
            "nit": nit,
            "actual": actual,
            "aseguradoras": [c.aseguradora for c in cotizaciones],
            "vehiculos_flota_comparable": len(placas),
            "advertencias": advertencias,
            "duration_seconds": round(monotonic() - started, 2),
        }
        return CotizadorOutput(
            content=contenido, filename=f"Consolidado_{ramo}_{nit}.xlsx", summary=summary,
        )
